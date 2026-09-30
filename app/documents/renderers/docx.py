"""Renderer DOCX — ``DocumentSpec`` → octets OOXML (ADR 0005, Phases 6-7).

Le renderer ne décide de **rien** sur le fond : il traduit une structure et des
tokens de style en un document Word professionnel. Tout ce que python-docx ne
sait pas exprimer est délégué à ``app.documents.docx_xml`` (couche isolée).

Pris en charge : page de couverture (titre, sous-titre, organisation, logo,
champs), titres 1→4, paragraphes (corps/chapeau/note), listes à puces et
numérotées imbriquées, tableaux (en-tête simple ou à étages, fusions
horizontales et verticales, trames par cellule, largeurs explicites, ligne de
total, en-tête répété en haut de page), images légendées, encadrés, citations,
blocs libellé/valeur, sauts de page, en-tête et pied de page distincts sur la
première page, numérotation de page.

Un tableau **sans en-tête déclaré** n'en reçoit aucun : le rendu n'invente
jamais de titre de colonne (une bande d'en-tête vide a déjà été produite par le
passé, et un tableau bancal est refusé par ``Tableau.grille``).

Règles tenues par ce module :

- **aucune valeur inventée** : un bloc de statut ``inconnu``/``manquant`` est
  rendu avec le marqueur ``(à compléter)``, jamais avec un blanc silencieux ;
- **aucune dépendance obligatoire** : python-docx est importé paresseusement via
  ``require_library`` — une librairie absente produit une ``GenerationError``
  explicite, jamais un fichier vide.

Limites documentées : les listes utilisent les styles Word ``List Bullet`` /
``List Number`` (deux listes numérotées successives peuvent poursuivre la
numérotation) ; le nombre de pages d'un DOCX n'est pas calculable sans moteur de
rendu — c'est le PDF dérivé qui donne la pagination réelle.
"""

from __future__ import annotations

import io
from typing import Any, NamedTuple

from app.core.errors import CarsoError, GenerationError, ValidationError
from app.documents import docx_xml as xml
from app.documents.assets import FournisseurImages, resoudre_image
from app.documents.renderers.commun import (
    alignement_cellule,
    couleur_cellule,
    fond_cellule,
    gras_cellule,
)
from app.documents.spec import (
    MARQUEUR_A_COMPLETER,
    MARQUEUR_A_CONFIRMER,
    MARQUEUR_RE,
    Bloc,
    Citation,
    Couverture,
    DocumentSpec,
    Encadre,
    Image,
    Liste,
    Paires,
    Paragraphe,
    SautDePage,
    Tableau,
    Titre,
)
from app.documents.styles import DocumentStyle, StyleTexte, couleur_rgb

__all__ = ["rendre_docx"]

_ALIGNEMENTS = {
    "gauche": "LEFT",
    "centre": "CENTER",
    "droite": "RIGHT",
    "justifie": "JUSTIFY",
}


class _Libs(NamedTuple):
    """Points d'entrée python-docx utilisés par le renderer (imports groupés)."""

    Document: Any
    Mm: Any
    Pt: Any
    RGBColor: Any
    alignements: Any
    orientations: Any


def _libs() -> _Libs:
    """Importe python-docx et ses unités (librairie optionnelle, extra ``documents``).

    Raises:
        GenerationError: python-docx absent.
    """
    try:
        from app.documents.extraction.base import require_library

        docx = require_library("docx")
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Mm, Pt, RGBColor
    except CarsoError as erreur:
        raise GenerationError(str(erreur)) from erreur
    return _Libs(
        Document=docx.Document,
        Mm=Mm,
        Pt=Pt,
        RGBColor=RGBColor,
        alignements=WD_ALIGN_PARAGRAPH,
        orientations=WD_ORIENT,
    )


def rendre_docx(spec: DocumentSpec, *, images: FournisseurImages | None = None) -> bytes:
    """Rend un ``DocumentSpec`` en octets DOCX.

    Args:
        spec: le document structuré à rendre.
        images: fournisseur d'images (``None`` si le document n'en contient pas).

    Raises:
        GenerationError: python-docx absent ou échec de construction.
        ValidationError: image référencée mais introuvable.
    """
    libs = _libs()
    style = spec.style_effectif()
    try:
        document = libs.Document()
        _configurer_styles_base(document, libs, style)
        section = document.sections[0]
        couverture = spec.couverture()
        _configurer_page(section, libs, spec, couverture is not None)
        _configurer_entete_pied(section, libs, spec, style)

        if couverture is not None:
            _rendre_couverture(document, libs, spec, style, couverture, images)
        for bloc in spec.blocs:
            _rendre_bloc(document, libs, bloc, spec, style, images)

        buffer = io.BytesIO()
        document.save(buffer)
    except (ValidationError, GenerationError):
        raise
    except Exception as erreur:  # pragma: no cover - défensif (bibliothèque tierce)
        raise GenerationError(f"Échec du rendu DOCX : {erreur}") from erreur
    return buffer.getvalue()


# --- Configuration du document -----------------------------------------------


def _configurer_styles_base(document: Any, libs: _Libs, style: DocumentStyle) -> None:
    """Aligne le style ``Normal`` sur les tokens du corps de texte."""
    normal = document.styles["Normal"]
    normal.font.name = style.corps.police
    normal.font.size = libs.Pt(style.corps.taille_pt)
    normal.font.color.rgb = libs.RGBColor(*couleur_rgb(style.corps.couleur))


def _configurer_page(
    section: Any, libs: _Libs, spec: DocumentSpec, avec_couverture: bool
) -> None:
    """Applique format, orientation, marges et règle de première page."""
    geometrie = spec.geometrie()
    section.page_width = libs.Mm(geometrie.largeur_mm)
    section.page_height = libs.Mm(geometrie.hauteur_mm)
    section.orientation = (
        libs.orientations.LANDSCAPE
        if geometrie.orientation == "paysage"
        else libs.orientations.PORTRAIT
    )
    section.left_margin = libs.Mm(geometrie.marges.gauche_mm)
    section.right_margin = libs.Mm(geometrie.marges.droite_mm)
    section.top_margin = libs.Mm(geometrie.marges.haut_mm)
    section.bottom_margin = libs.Mm(geometrie.marges.bas_mm)
    # Une couverture ne porte ni en-tête ni pied de page : elle reste sobre.
    section.different_first_page_header_footer = avec_couverture


def _configurer_entete_pied(
    section: Any, libs: _Libs, spec: DocumentSpec, style: DocumentStyle
) -> None:
    """Écrit l'en-tête et le pied de page (texte, filet, numéro de page)."""
    mention = spec.mise_en_page.mention_proposition
    texte_entete = spec.mise_en_page.entete_texte or style.entete.texte or mention or ""
    entete = section.header.paragraphs[0]
    entete.text = ""
    if texte_entete:
        run = entete.add_run(texte_entete)
        _styler_run(run, libs, style.entete.comme_style_texte())
    _aligner(entete, libs, style.entete.alignement)
    xml.espacement_paragraphe(entete, avant_pt=0.0, apres_pt=0.0)

    pied = section.footer.paragraphs[0]
    pied.text = ""
    texte_pied = spec.mise_en_page.pied_texte or style.pied.texte or ""
    if texte_pied:
        pied.add_run(texte_pied)
    numero = (
        spec.mise_en_page.numero_de_page
        if spec.mise_en_page.numero_de_page is not None
        else style.pied.numero_page
    )
    if numero:
        if texte_pied:
            pied.add_run("  —  ")
        xml.champ_numero_page(pied)
    for run in pied.runs:
        _styler_run(run, libs, style.pied.comme_style_texte())
    _aligner(pied, libs, style.pied.alignement)
    if style.pied.filet:
        xml.filet_haut_paragraphe(pied, couleur=style.pied.couleur, taille_pt=0.5)
    xml.espacement_paragraphe(pied, avant_pt=2.0, apres_pt=0.0)


def _rendre_couverture(
    document: Any,
    libs: _Libs,
    spec: DocumentSpec,
    style: DocumentStyle,
    couverture: Couverture,
    images: FournisseurImages | None,
) -> None:
    """Rend la page de couverture, puis un saut de page."""
    tokens = style.couverture
    document.add_paragraph()
    document.add_paragraph()
    if couverture.logo_document_id is not None or couverture.logo_chemin:
        _inserer_image(
            document,
            libs,
            Image(
                document_id=couverture.logo_document_id,
                chemin=couverture.logo_chemin,
                largeur_mm=tokens.logo_largeur_mm,
            ),
            images,
            centrer=True,
        )
    titre = couverture.titre or spec.metadata.titre
    _ajouter_paragraphe(
        document,
        libs,
        titre,
        tokens=StyleTexte(
            police=style.titre.police,
            taille_pt=tokens.titre_taille_pt,
            gras=True,
            couleur=tokens.couleur_accent,
            alignement=tokens.alignement,
            espace_apres_pt=6.0,
        ),
    )
    sous_titre = couverture.sous_titre or spec.metadata.sous_titre
    if sous_titre:
        _ajouter_paragraphe(
            document,
            libs,
            sous_titre,
            tokens=StyleTexte(
                police=style.corps.police,
                taille_pt=tokens.sous_titre_taille_pt,
                couleur=style.corps.couleur,
                alignement=tokens.alignement,
                espace_apres_pt=10.0,
            ),
        )
    if tokens.bandeau:
        _ajouter_paragraphe(
            document,
            libs,
            "—" * 24,
            tokens=StyleTexte(
                taille_pt=tokens.sous_titre_taille_pt,
                couleur=tokens.couleur_accent,
                alignement=tokens.alignement,
                espace_apres_pt=10.0,
            ),
        )
    organisation = couverture.organisation or spec.metadata.organisation
    if organisation:
        _ajouter_paragraphe(
            document,
            libs,
            organisation,
            tokens=StyleTexte(
                taille_pt=style.corps.taille_pt + 1,
                gras=True,
                alignement=tokens.alignement,
                espace_apres_pt=4.0,
            ),
        )
    for paire in couverture.champs:
        _ajouter_paragraphe(
            document,
            libs,
            f"{paire.libelle} : {paire.valeur}",
            tokens=StyleTexte(
                taille_pt=style.corps.taille_pt,
                couleur=style.corps.couleur,
                alignement=tokens.alignement,
                espace_apres_pt=2.0,
            ),
        )
    if couverture.date_ligne:
        _ajouter_paragraphe(
            document,
            libs,
            couverture.date_ligne,
            tokens=StyleTexte(
                taille_pt=style.corps.taille_pt,
                italique=True,
                alignement=tokens.alignement,
                espace_apres_pt=2.0,
            ),
        )
    if spec.mise_en_page.mention_proposition:
        _ajouter_paragraphe(
            document,
            libs,
            spec.mise_en_page.mention_proposition,
            tokens=StyleTexte(
                taille_pt=8.5,
                italique=True,
                couleur="#6B7280",
                alignement=tokens.alignement,
                espace_avant_pt=24.0,
                espace_apres_pt=0.0,
            ),
        )
    document.add_page_break()


# --- Rendu des blocs ---------------------------------------------------------


def _rendre_bloc(
    document: Any,
    libs: _Libs,
    bloc: Bloc,
    spec: DocumentSpec,
    style: DocumentStyle,
    images: FournisseurImages | None,
) -> None:
    """Traduit un bloc du spec en contenu DOCX."""
    if isinstance(bloc, Titre):
        paragraphe = document.add_paragraph(bloc.texte, style=f"Heading {bloc.niveau}")
        tokens = style.titre_niveau(bloc.niveau)
        _appliquer(paragraphe, libs, tokens)
        _ajouter_marqueur(
            paragraphe,
            libs,
            tokens,
            bloc.statut,
            deja_marque=MARQUEUR_RE.search(bloc.texte),
        )
        return

    if isinstance(bloc, Paragraphe):
        tokens = style.citation if bloc.role == "note" else style.corps
        if bloc.role == "chapeau":
            tokens = tokens.model_copy(
                update={
                    "taille_pt": tokens.taille_pt + 1.0,
                    "espace_apres_pt": tokens.espace_apres_pt + 4.0,
                }
            )
        paragraphe = _ajouter_paragraphe(document, libs, bloc.texte, tokens=tokens)
        _ajouter_marqueur(
            paragraphe,
            libs,
            tokens,
            bloc.statut,
            deja_marque=MARQUEUR_RE.search(bloc.texte),
        )
        return

    if isinstance(bloc, Liste):
        _rendre_liste(document, libs, bloc, style)
        return

    if isinstance(bloc, Tableau):
        _rendre_tableau(document, libs, bloc, spec, style)
        return

    if isinstance(bloc, Image):
        _inserer_image(document, libs, bloc, images, centrer=True)
        return

    if isinstance(bloc, Encadre):
        _rendre_encadre(document, libs, bloc, style)
        return

    if isinstance(bloc, Citation):
        paragraphe = _ajouter_paragraphe(document, libs, f"« {bloc.texte} »", tokens=style.citation)
        _ajouter_marqueur(paragraphe, libs, style.citation, bloc.statut, deja_marque=False)
        if bloc.attribution:
            _ajouter_paragraphe(
                document,
                libs,
                f"— {bloc.attribution}",
                tokens=style.citation.model_copy(
                    update={"alignement": "droite", "espace_apres_pt": 8.0}
                ),
            )
        return

    if isinstance(bloc, Paires):
        _rendre_paires(document, libs, bloc, style)
        return

    if isinstance(bloc, SautDePage):
        document.add_page_break()
        return

    raise GenerationError(f"Type de bloc non rendable : {type(bloc).__name__}")


def _rendre_liste(document: Any, libs: _Libs, bloc: Liste, style: DocumentStyle) -> None:
    """Liste à puces ou numérotée, indentée selon sa profondeur."""
    modele = "List Number" if bloc.numerotee else "List Bullet"
    for item in bloc.items:
        paragraphe = document.add_paragraph(item, style=modele)
        _appliquer(paragraphe, libs, style.liste)
        if bloc.profondeur:
            paragraphe.paragraph_format.left_indent = libs.Mm(
                6.0 + 4.0 * bloc.profondeur
            )
        _ajouter_marqueur(
            paragraphe, libs, style.liste, bloc.statut, deja_marque=MARQUEUR_RE.search(item)
        )


def _rendre_tableau(
    document: Any, libs: _Libs, bloc: Tableau, spec: DocumentSpec, style: DocumentStyle
) -> None:
    """Tableau : en-tête stylé et répété, largeurs explicites, bandes alternées."""
    tokens = style.tableau
    if bloc.titre_tableau:
        _ajouter_paragraphe(
            document,
            libs,
            bloc.titre_tableau,
            tokens=style.corps.model_copy(
                update={"gras": True, "espace_avant_pt": 8.0, "espace_apres_pt": 3.0}
            ),
        )
    grille = bloc.grille()
    nb_colonnes = bloc.largeur_grille()
    nb_entete = bloc.nb_lignes_entete()

    table = document.add_table(rows=len(grille), cols=nb_colonnes)
    table.style = "Table Grid"

    # 1) Fusions d'abord : une cellule fusionnée ne contient alors qu'un seul
    #    paragraphe (sinon les paragraphes vides des cellules absorbées restent).
    for rangees in grille:
        for placee in rangees:
            if placee.colonnes == 1 and placee.lignes == 1:
                continue
            table.cell(placee.ligne, placee.colonne).merge(
                table.cell(placee.ligne + placee.lignes - 1, placee.colonne + placee.colonnes - 1)
            )

    # 2) Contenu et style, aux seules positions d'ancrage.
    for index_ligne, rangees in enumerate(grille):
        ligne_docx = table.rows[index_ligne]
        en_entete = index_ligne < nb_entete
        index_donnee = index_ligne - nb_entete
        total = bloc.total_ligne is not None and index_donnee == bloc.total_ligne
        alternance = bool(tokens.bandes_alternees and index_donnee % 2 == 1 and not total)
        if en_entete and tokens.repeter_entete and bloc.repeter_entete:
            xml.repetition_entete(ligne_docx)
        for placee in rangees:
            cellule = table.cell(placee.ligne, placee.colonne)
            cellule.text = ""
            run = cellule.paragraphs[0].add_run(placee.cellule.texte)
            _styler_run(
                run,
                libs,
                StyleTexte(
                    police=style.corps.police,
                    taille_pt=placee.cellule.taille_pt or tokens.taille_pt,
                    gras=gras_cellule(
                        placee.cellule, en_entete=en_entete, total=total, tokens=tokens
                    ),
                    couleur=couleur_cellule(
                        placee.cellule, en_entete=en_entete, style=style, tokens=tokens
                    ),
                ),
            )
            fond = fond_cellule(
                placee.cellule,
                en_entete=en_entete,
                total=total,
                alternance=alternance,
                tokens=tokens,
            )
            if fond:
                xml.ombrer_cellule(cellule, fond)
            _aligner_cellule(cellule, libs, bloc, placee)

    xml.mise_en_page_fixe(table)
    xml.largeurs_colonnes(table, bloc.largeurs_mm or _largeurs_par_defaut(spec, nb_colonnes))
    if tokens.bordures:
        xml.bordures_tableau(
            table, couleur=tokens.bordure_couleur, taille_pt=tokens.bordure_taille_pt
        )
    document.add_paragraph()


def _rendre_encadre(document: Any, libs: _Libs, bloc: Encadre, style: DocumentStyle) -> None:
    """Encadré : tableau d'une cellule tramé et bordé (robuste Word/LibreOffice)."""
    tokens = style.encadre(bloc.genre)
    table = document.add_table(rows=1, cols=1)
    cellule = table.cell(0, 0)
    cellule.text = ""
    paragraphe = cellule.paragraphs[0]
    libelle = bloc.titre or tokens.libelle
    if libelle:
        run = paragraphe.add_run(f"{libelle} — ")
        _styler_run(
            run,
            libs,
            StyleTexte(
                police=style.corps.police,
                taille_pt=style.corps.taille_pt,
                gras=True,
                couleur=tokens.couleur_titre,
            ),
        )
    run = paragraphe.add_run(bloc.texte)
    _styler_run(
        run,
        libs,
        StyleTexte(
            police=style.corps.police,
            taille_pt=style.corps.taille_pt,
            couleur=tokens.couleur_texte,
        ),
    )
    xml.ombrer_cellule(cellule, tokens.fond)
    xml.bordures_tableau(table, couleur=tokens.bordure, taille_pt=0.75)
    document.add_paragraph()


def _rendre_paires(document: Any, libs: _Libs, bloc: Paires, style: DocumentStyle) -> None:
    """Bloc libellé/valeur : tableau sans filets, libellés en gras."""
    nb_colonnes = bloc.colonnes * 2
    table = document.add_table(rows=0, cols=nb_colonnes)
    paires = list(bloc.paires)
    for debut in range(0, len(paires), bloc.colonnes):
        ligne = table.add_row()
        for index in range(bloc.colonnes):
            position = debut + index
            cellule_libelle = ligne.cells[index * 2]
            cellule_valeur = ligne.cells[index * 2 + 1]
            cellule_libelle.text = ""
            cellule_valeur.text = ""
            if position >= len(paires):
                continue
            paire = paires[position]
            run = cellule_libelle.paragraphs[0].add_run(f"{paire.libelle} :")
            _styler_run(
                run,
                libs,
                StyleTexte(
                    police=style.corps.police,
                    taille_pt=style.tableau.taille_pt,
                    gras=True,
                    couleur=style.corps.couleur,
                ),
            )
            run = cellule_valeur.paragraphs[0].add_run(paire.valeur)
            _styler_run(
                run,
                libs,
                StyleTexte(
                    police=style.corps.police,
                    taille_pt=style.tableau.taille_pt,
                    couleur=style.corps.couleur,
                ),
            )
    document.add_paragraph()


# --- Blocs simples ----------------------------------------------------------


def _inserer_image(
    document: Any,
    libs: _Libs,
    bloc: Image,
    images: FournisseurImages | None,
    *,
    centrer: bool,
) -> None:
    """Insère une image (ou lève une erreur explicite si elle est introuvable)."""
    contenu = resoudre_image(bloc, images)
    if contenu is None:
        raise ValidationError(
            "Image référencée sans fournisseur d'images",
            details={"document_id": str(bloc.document_id or ""), "chemin": bloc.chemin},
        )
    octets, nom = contenu
    try:
        paragraphe = document.add_paragraph()
        paragraphe.add_run().add_picture(io.BytesIO(octets), width=libs.Mm(bloc.largeur_mm))
    except Exception as erreur:
        raise ValidationError(f"Image illisible : {nom}") from erreur
    if centrer:
        _aligner(paragraphe, libs, "centre")
    if bloc.legende:
        _ajouter_paragraphe(
            document,
            libs,
            bloc.legende,
            tokens=StyleTexte(
                taille_pt=8.5,
                italique=True,
                couleur="#6B7280",
                alignement="centre",
                espace_apres_pt=8.0,
            ),
        )


def _ajouter_paragraphe(
    document: Any, libs: _Libs, texte: str, tokens: StyleTexte | None = None
) -> Any:
    """Ajoute un paragraphe stylé à partir des tokens."""
    paragraphe = document.add_paragraph()
    paragraphe.add_run(texte)
    _appliquer(paragraphe, libs, tokens)
    return paragraphe


def _ajouter_marqueur(
    paragraphe: Any,
    libs: _Libs,
    tokens: StyleTexte | None,
    statut: str,
    *,
    deja_marque: object,
) -> None:
    """Ajoute le marqueur visible d'une donnée absente (``(à compléter)``)."""
    if deja_marque:
        return
    marqueur = _marqueur_du_statut(statut)
    if not marqueur:
        return
    run = paragraphe.add_run(f" {marqueur}")
    if tokens is not None:
        _styler_run(run, libs, tokens)


def _marqueur_du_statut(statut: str) -> str:
    """Marqueur affiché pour un statut de renseignement non satisfait."""
    if statut in ("inconnu", "manquant"):
        return MARQUEUR_A_COMPLETER
    if statut == "a_confirmer":
        return MARQUEUR_A_CONFIRMER
    return ""


def _appliquer(
    paragraphe: Any, libs: _Libs, tokens: StyleTexte | None, run: Any | None = None
) -> None:
    """Applique un rôle typographique (alignement, espacement, police)."""
    if tokens is None:
        return
    _aligner(paragraphe, libs, tokens.alignement)
    xml.espacement_paragraphe(
        paragraphe,
        avant_pt=tokens.espace_avant_pt,
        apres_pt=tokens.espace_apres_pt,
        interligne=tokens.interligne,
    )
    cible = run if run is not None else (paragraphe.runs[0] if paragraphe.runs else None)
    if cible is not None:
        _styler_run(cible, libs, tokens)


def _styler_run(run: Any, libs: _Libs, tokens: StyleTexte) -> None:
    """Applique police, taille, graisse, italique et couleur à un run."""
    run.font.name = tokens.police
    run.font.size = libs.Pt(tokens.taille_pt)
    run.font.bold = tokens.gras
    run.font.italic = tokens.italique
    run.font.color.rgb = libs.RGBColor(*couleur_rgb(tokens.couleur))


def _aligner(paragraphe: Any, libs: _Libs, alignement: str) -> None:
    """Applique un alignement de paragraphe."""
    paragraphe.alignment = getattr(libs.alignements, _ALIGNEMENTS[alignement])


def _aligner_cellule(cellule: Any, libs: _Libs, bloc: Tableau, placee: Any) -> None:
    """Applique l'alignement de la cellule, sinon celui de sa colonne de départ."""
    alignement = alignement_cellule(placee.cellule, bloc.alignements, placee.colonne)
    if alignement is None:
        return
    for paragraphe in cellule.paragraphs:
        _aligner(paragraphe, libs, alignement)


def _largeurs_par_defaut(spec: DocumentSpec, nb_colonnes: int) -> list[float]:
    """Répartit la largeur utile quand aucune largeur n'est fournie."""
    utile = spec.geometrie().largeur_utile_mm()
    return [round(utile / nb_colonnes, 2)] * nb_colonnes
