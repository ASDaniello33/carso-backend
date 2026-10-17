"""Rendu Excel professionnel (``openpyxl``) — les tableaux du corps HTML.

Le classeur est construit **depuis le même corps HTML** que le PDF et le DOCX :
chaque ``table.data`` devient une feuille stylée thème CARSO (en-tête figé,
colonnes dimensionnées, lignes alternées), les ``table.kpi`` alimentent une
feuille « Chiffres clés ». La mise en page est décidée ici, jamais par l'agent ;
sans aucun tableau, le classeur est refusé explicitement (jamais un .xlsx vide
présenté comme un livrable).
"""

from __future__ import annotations

import re
from pathlib import Path

from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.errors import ValidationError

from .themes import BORDURE, MUTED, TEXTE, Theme, get_theme

__all__ = ["rendre_xlsx"]

_TIROIR_MAX = 31  # limite de nom de feuille Excel
_LARGEUR_MIN, _LARGEUR_MAX = 10, 60


# --------------------------------------------------------------------------- lecture
def _texte(el: Tag | None) -> str:
    if el is None:
        return ""
    return " ".join(el.get_text(" ", strip=True).split())


def _classes(el: Tag) -> set[str]:
    valeur: object = el.get("class") or []
    if isinstance(valeur, str):
        return set(valeur.split())
    return {str(classe) for classe in valeur}


def _colonnes_numeriques(table: Tag) -> set[int]:
    numeriques: set[int] = set()
    ligne = table.find("tr")
    if ligne is None:
        return numeriques
    for index, cellule in enumerate(ligne.find_all(["td", "th"], recursive=False)):
        if "num" in _classes(cellule):
            numeriques.add(index)
    return numeriques


def _lignes_tableau(table: Tag) -> tuple[list[str], list[list[str]]]:
    entetes: list[str] = []
    lignes: list[list[str]] = []
    for tr in table.find_all("tr"):
        if tr.find_parent("table") is not table:
            continue
        cellules = tr.find_all(["td", "th"], recursive=False)
        if not cellules:
            continue
        valeurs = [_texte(cellule) for cellule in cellules]
        parent = tr.parent.name if tr.parent is not None else ""
        if parent == "thead" or (not entetes and cellules[0].name == "th"):
            entetes = valeurs
            continue
        if parent == "tfoot":
            continue  # un total déjà présent dans les lignes reste lisible
        lignes.append(valeurs)
    if not entetes and lignes:
        entetes = [f"Colonne {i + 1}" for i in range(len(lignes[0]))]
    largeur = len(entetes)
    lignes = [(ligne + [""] * largeur)[:largeur] for ligne in lignes]
    return entetes, lignes


def _nom_feuille(racine: Tag, table: Tag, deja_pris: set[str], rang: int) -> str:
    """Nom de feuille : le titre le plus proche AVANT le tableau, sinon Tableau n."""
    titre = ""
    for frere in racine.descendants:
        if frere is table:
            break
        if isinstance(frere, Tag) and frere.name in ("h1", "h2", "h3"):
            texte = _texte(frere)
            if texte:
                titre = texte
    brut = titre or f"Tableau {rang + 1}"
    nom = re.sub(r"[\[\]:*?/\\]", " ", brut).strip()[:_TIROIR_MAX] or f"Tableau {rang + 1}"
    base = nom
    suffixe = 2
    while nom.lower() in deja_pris:
        nom = f"{base[:_TIROIR_MAX - 3]} ({suffixe})"
        suffixe += 1
    deja_pris.add(nom.lower())
    return nom


def _kpis_classeur(racine: Tag) -> list[tuple[str, str]]:
    kpis: list[tuple[str, str]] = []
    for table in racine.find_all("table"):
        if "kpi" not in _classes(table):
            continue
        for cellule in table.find_all(["td", "th"]):
            valeur = cellule.find(class_="kpi-value")
            label = cellule.find(class_="kpi-label")
            if valeur is not None:
                kpis.append((_texte(label), _texte(valeur)))
            else:
                texte = _texte(cellule)
                if texte:
                    kpis.append(("", texte))
    return kpis


def _tables_donnees(racine: Tag) -> list[tuple[str, Tag, set[int], list[str], list[list[str]]]]:
    tables: list[tuple[str, Tag, set[int], list[str], list[list[str]]]] = []
    deja_pris: set[str] = set()
    for rang, table in enumerate(racine.find_all("table")):
        if "data" not in _classes(table):
            continue
        entetes, lignes = _lignes_tableau(table)
        if not entetes and not lignes:
            continue
        nom = _nom_feuille(racine, table, deja_pris, rang)
        tables.append((nom, table, _colonnes_numeriques(table), entetes, lignes))
    return tables


# --------------------------------------------------------------------------- style
def _bord_basse(couleur: str) -> Border:
    cote = Side(style="thin", color=couleur.lstrip("#").upper())
    return Border(bottom=cote)


def _feuille_tableau(
    wb: Workbook,
    nom: str,
    t: Theme,
    entetes: list[str],
    lignes: list[list[str]],
    numeriques: set[int],
    titre_document: str,
) -> None:
    ws = wb.create_sheet(title=nom[:_TIROIR_MAX])
    largeur = max(len(entetes), 1)
    # Titre du document, fusionné sur la largeur du tableau.
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=largeur)
    cellule_titre = ws.cell(row=1, column=1, value=titre_document)
    cellule_titre.font = Font(bold=True, size=13, color=t.primary.lstrip("#").upper())
    cellule_titre.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    # En-têtes : fond primaire, blanc, gras.
    remplir_en_tete = PatternFill("solid", start_color=t.primary.lstrip("#").upper())
    for colonne, libelle in enumerate(entetes, start=1):
        cellule = ws.cell(row=2, column=colonne, value=libelle)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = remplir_en_tete
        cellule.alignment = Alignment(
            horizontal="right" if (colonne - 1) in numeriques else "left",
            vertical="center",
            wrap_text=False,
        )
        cellule.border = _bord_basse(t.primary)

    alternance = PatternFill("solid", start_color=t.soft.lstrip("#").upper())
    for rang, ligne in enumerate(lignes, start=3):
        pair = (rang - 3) % 2 == 1
        for colonne, valeur in enumerate(ligne, start=1):
            cellule = ws.cell(row=rang, column=colonne, value=valeur)
            cellule.font = Font(color=TEXTE.lstrip("#").upper())
            cellule.alignment = Alignment(
                horizontal="right" if (colonne - 1) in numeriques else "left",
                vertical="center",
            )
            cellule.border = _bord_basse(BORDURE)
            if pair:
                cellule.fill = alternance

    # Largeurs de colonnes : le plus long contenu, borné.
    for colonne in range(1, largeur + 1):
        contenus = [
            len(entetes[colonne - 1]),
            *(len(ligne[colonne - 1]) for ligne in lignes),
        ]
        ws.column_dimensions[get_column_letter(colonne)].width = max(
            _LARGEUR_MIN, min(_LARGEUR_MAX, max(contenus, default=10) + 2)
        )

    ws.freeze_panes = "A3"  # titre + en-têtes restent visibles
    ws.auto_filter.ref = (
        f"A2:{get_column_letter(largeur)}{max(2, 2 + len(lignes))}"
    )


def _feuille_kpis(wb: Workbook, t: Theme, kpis: list[tuple[str, str]], titre_document: str) -> None:
    ws = wb.create_sheet(title="Chiffres clés")
    ws.merge_cells("A1:B1")
    cellule_titre = ws.cell(row=1, column=1, value=titre_document)
    cellule_titre.font = Font(bold=True, size=13, color=t.primary.lstrip("#").upper())
    ws.row_dimensions[1].height = 26
    for rang, (label, valeur) in enumerate(kpis, start=3):
        cellule_label = ws.cell(row=rang, column=1, value=(label or "Indicateur").upper())
        cellule_label.font = Font(size=9, color=MUTED.lstrip("#").upper())
        cellule_valeur = ws.cell(row=rang, column=2, value=valeur)
        cellule_valeur.font = Font(bold=True, size=12, color=t.primary.lstrip("#").upper())
        cellule_valeur.alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18


# --------------------------------------------------------------------------- point d'entrée
def rendre_xlsx(
    corps_html: str,
    chemin: str | Path,
    *,
    theme: str | Theme = "corporate",
    titre: str = "",
    auteur: str = "",
) -> Path:
    """Écrit le classeur et retourne son chemin.

    Raises:
        ValidationError: corps vide, thème inconnu, ou aucun tableau
            (``table.data`` / ``table.kpi``) : un .xlsx sans données n'est pas
            un livrable.
    """
    if not corps_html or not corps_html.strip():
        raise ValidationError("contenu_html est vide : rien à mettre en classeur.")
    t = get_theme(theme)

    racine = BeautifulSoup(corps_html, "html.parser")
    tables = _tables_donnees(racine)
    kpis = _kpis_classeur(racine)
    if not tables and not kpis:
        raise ValidationError(
            "Un classeur Excel attend des tableaux : aucun table.data (ni "
            "table.kpi) dans le contenu. Écris au moins un tableau de données "
            "avec <thead>, ou choisis le format pdf/docx pour un document texte.",
            details={"attendus": ["table.data", "table.kpi"]},
        )

    wb = Workbook()
    wb.remove(wb.active)  # la feuille par défaut : on ne garde que les nôtres
    for nom, _table, numeriques, entetes, lignes in tables:
        _feuille_tableau(wb, nom, t, entetes, lignes, numeriques, titre)
    if kpis:
        _feuille_kpis(wb, t, kpis, titre)

    wb.properties.title = titre or "Classeur"
    wb.properties.creator = auteur

    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(chemin))
    return chemin
