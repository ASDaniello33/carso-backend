"""ADR 0005 — renderers et couche XML DOCX (Phases 6 et 8).

Les tests réinspectent les **fichiers produits** avec les mêmes librairies que la
production (python-docx, pypdf) : ils vérifient donc le rendu réel, pas l'appel
de fonction. La couche XML isolée (``docx_xml``) est testée séparément : c'est
elle qui porte les cas que python-docx n'expose pas (répétition d'en-tête,
trames, largeurs fixes, champ de pagination).
"""

from __future__ import annotations

import io

import pytest
from document_fixtures import build_png

from app.core.errors import GenerationError, ValidationError
from app.documents.assets import ImagesMemoire
from app.documents.docx_xml import (
    bordures_tableau,
    champ_numero_page,
    espacement_paragraphe,
    filet_haut_paragraphe,
    largeurs_colonnes,
    mise_en_page_fixe,
    ombrer_cellule,
    repetition_entete,
)
from app.documents.renderers import FORMATS_RENDUS, extension_et_mime, rendre
from app.documents.spec import (
    Citation,
    Couverture,
    DocumentSpec,
    Encadre,
    Image,
    Liste,
    Metadonnees,
    MiseEnPage,
    Paire,
    Paires,
    Paragraphe,
    SautDePage,
    Tableau,
    Titre,
)
from app.documents.styles import CARSO_DEFAUT, Marges, StyleCouverture

_LOGO = "offres/1/logo.png"


def _images() -> ImagesMemoire:
    """Fournisseur d'images de test : le logo du spec complet."""
    return ImagesMemoire({_LOGO: build_png(8, 6)})


def _spec_complet(*, paysage: bool = False) -> DocumentSpec:
    """Document complet : couverture, titres, listes, tableau, image, encadré."""
    return DocumentSpec(
        metadata=Metadonnees(
            titre="Offre technique — Lot 3",
            sous_titre="Appel AAP-2026-07",
            organisation="CARSO",
            reference="OT-2026-014",
            auteur="AgentGenerateurOffre",
        ),
        mise_en_page=MiseEnPage(orientation="paysage" if paysage else None),
        blocs=[
            Titre(niveau=1, texte="Contexte"),
            Paragraphe(texte="Zone couverte : Rabat et Salé."),
            Liste(items=["Cadrage", "Exécution"]),
            Liste(items=["Phase 1"], numerotee=True),
            Tableau(
                titre_tableau="Budget prévisionnel",
                entete=["Poste", "Montant"],
                lignes=[["Formateurs", "12 000"], ["Total", "12 000"]],
                largeurs_mm=[110, 62],
                alignements=["gauche", "droite"],
                total_ligne=1,
            ),
            Encadre(genre="avertissement", texte="Le calendrier reste à valider."),
            Citation(texte="Extrait du cahier des charges", attribution="Organisation"),
            Paires(paires=[Paire(libelle="Lieu", valeur="Rabat")]),
            Image(chemin=_LOGO, largeur_mm=25, legende="Logo client"),
            SautDePage(),
            Titre(niveau=2, texte="Annexes"),
            Paragraphe(texte="Effectif : (à compléter)"),
        ],
    )


class TestDocxRenderer:
    def test_structure_et_couverture(self) -> None:
        import docx

        document = docx.Document(
            io.BytesIO(rendre(_spec_complet(), format="docx", images=_images()))
        )
        styles = [p.style.name for p in document.paragraphs]
        textes = "\n".join(p.text for p in document.paragraphs)

        assert "Heading 1" in styles and "Heading 2" in styles
        assert "Offre technique — Lot 3" in textes
        assert "Document proposé" in textes
        assert "List Bullet" in styles and "List Number" in styles
        assert "Extrait du cahier des charges" in textes
        encadres = [t for t in document.tables if len(t.columns) == 1]
        assert "Le calendrier reste à valider." in encadres[0].cell(0, 0).text
        assert document.inline_shapes, "l'image doit être intégrée au DOCX"

    def test_page_format_et_marges(self) -> None:
        import docx

        portrait = docx.Document(
            io.BytesIO(rendre(_spec_complet(), format="docx", images=_images()))
        )
        section = portrait.sections[0]
        assert round(section.page_width.mm) == 210
        assert round(section.page_height.mm) == 297
        assert round(section.left_margin.mm) == 20
        assert section.different_first_page_header_footer is True

        paysage = docx.Document(
            io.BytesIO(rendre(_spec_complet(paysage=True), format="docx", images=_images()))
        )
        assert paysage.sections[0].page_width > paysage.sections[0].page_height

    def test_entete_pied_et_numero_de_page(self) -> None:
        import docx

        document = docx.Document(
            io.BytesIO(rendre(_spec_complet(), format="docx", images=_images()))
        )
        section = document.sections[0]
        assert "Document proposé" in section.header.paragraphs[0].text
        assert "PAGE" in section.footer.paragraphs[0]._p.xml

    def test_tableau_largeurs_entete_repete(self) -> None:
        import docx
        from docx.oxml.ns import qn

        document = docx.Document(
            io.BytesIO(rendre(_spec_complet(), format="docx", images=_images()))
        )
        tableau = next(t for t in document.tables if len(t.columns) == 2)
        assert tableau.rows[0]._tr.find(qn("w:trPr")) is not None
        assert tableau.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
        assert tableau.rows[0].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:shd")) is not None
        grille = tableau._tbl.find(qn("w:tblGrid"))
        largeurs = [int(col.get(qn("w:w"))) for col in grille.findall(qn("w:gridCol"))]
        assert largeurs == [6236, 3515]  # 110 mm et 62 mm en twips

    def test_marqueur_donnee_manquante(self) -> None:
        """Aucune donnée absente rendue en silence : le marqueur est visible."""
        import docx

        document = docx.Document(
            io.BytesIO(rendre(_spec_complet(), format="docx", images=_images()))
        )
        texte = "\n".join(p.text for p in document.paragraphs)
        assert "(à compléter)" in texte

    def test_image_manquante_refusee(self) -> None:
        spec = DocumentSpec(
            metadata=Metadonnees(titre="Avec image"),
            blocs=[Image(chemin="offres/1/absent.png", largeur_mm=20)],
        )
        with pytest.raises(ValidationError, match="fournisseur"):
            rendre(spec, format="docx")

    def test_format_legacy_respecte(self) -> None:
        """Les adaptateurs historiques gardent titre + corps sur la même page."""
        from app.documents.generation import build_docx

        document_bytes = build_docx(
            titre="Fiche technique",
            paragraphes=["Mission M-2026-9"],
            tableau=[["Nom", "Rôle"], ["Alaoui", "formateur"]],
        )
        import docx

        document = docx.Document(io.BytesIO(document_bytes))
        assert [p.text for p in document.paragraphs][0] == "Fiche technique"
        assert document.tables[0].cell(1, 0).text == "Alaoui"


class TestPdfRenderer:
    def test_texte_extractible_et_pagination(self) -> None:
        from pypdf import PdfReader

        octets = rendre(_spec_complet(), format="pdf", images=_images())
        lecteur = PdfReader(io.BytesIO(octets))
        assert len(lecteur.pages) >= 2
        contenu = "\n".join(page.extract_text() or "" for page in lecteur.pages)
        assert "Offre technique" in contenu
        assert "Formateurs" in contenu and "12 000" in contenu
        assert "Cadrage" in contenu
        assert "Le calendrier reste à valider." in contenu
        assert "/ 2" in contenu or "Page 2" in contenu

    def test_metadonnees(self) -> None:
        from pypdf import PdfReader

        lecteur = PdfReader(io.BytesIO(rendre(_spec_complet(), format="pdf", images=_images())))
        assert lecteur.metadata is not None
        assert "Offre technique" in (lecteur.metadata.title or "")

    def test_translitteration_typographique(self) -> None:
        """Un « € » ou un tiret long ne casse jamais le rendu."""
        from app.documents.renderers.pdf import texte_pdf

        assert texte_pdf("Montant — 1 500 €") == "Montant - 1 500 EUR"
        assert texte_pdf("• puce") == "- puce"
        assert texte_pdf("œuf") == "oeuf"
        assert texte_pdf("\u2603") == "?"

        from pypdf import PdfReader

        spec = DocumentSpec(
            metadata=Metadonnees(titre="Budget"),
            mise_en_page=MiseEnPage(couverture=Couverture(active=False)),
            blocs=[Paragraphe(texte="Coût total : 12 000 € — garanti.")],
        )
        octets = rendre(spec, format="pdf")
        contenu = PdfReader(io.BytesIO(octets)).pages[0].extract_text() or ""
        assert "12 000 EUR" in contenu

    def test_paysage_et_geometrie(self) -> None:
        from pypdf import PdfReader

        paysage = _spec_complet(paysage=True)
        lecteur = PdfReader(io.BytesIO(rendre(paysage, format="pdf", images=_images())))
        boite = lecteur.pages[0].mediabox
        assert float(boite.width) > float(boite.height)

    def test_sans_couverture(self) -> None:
        from pypdf import PdfReader

        spec = DocumentSpec(
            metadata=Metadonnees(titre="Note interne"),
            mise_en_page=MiseEnPage(
                couverture=Couverture(active=False), mention_proposition=None
            ),
            blocs=[Paragraphe(texte="Une page.")],
        )
        lecteur = PdfReader(io.BytesIO(rendre(spec, format="pdf")))
        assert len(lecteur.pages) == 1


class TestHtmlRenderer:
    def test_structure_et_echappement(self) -> None:
        from app.documents.renderers.html import rendre_html_texte

        spec = DocumentSpec(
            metadata=Metadonnees(titre="Offre <script>alert(1)</script>"),
            blocs=[Paragraphe(texte="Risque : <b>injection</b>")],
        )
        html = rendre_html_texte(spec)
        assert "&lt;script&gt;" in html and "<script>" not in html
        assert "&lt;b&gt;injection&lt;/b&gt;" in html
        assert "<h2>Contexte</h2>" not in html  # aucun titre inventé

    def test_tableau_et_image_en_data_uri(self) -> None:
        from app.documents.assets import ImagesMemoire
        from app.documents.renderers.html import rendre_html_texte

        images = ImagesMemoire({"offres/1/logo.png": build_png()})
        spec = DocumentSpec(
            metadata=Metadonnees(titre="Offre"),
            blocs=[
                Tableau(entete=["A", "B"], lignes=[["1", "2"]], largeurs_mm=[80, 80]),
                Image(chemin="offres/1/logo.png", largeur_mm=20),
            ],
        )
        html = rendre_html_texte(spec, images=images)
        assert "<thead>" in html and "<colgroup>" in html
        assert "data:image/png;base64," in html

    def test_tokens_de_style_dans_le_css(self) -> None:
        from app.documents.renderers.html import rendre_html_texte

        spec = DocumentSpec(
            metadata=Metadonnees(titre="Offre"),
            mise_en_page=MiseEnPage(marges=Marges(gauche_mm=30, droite_mm=20)),
        )
        html = rendre_html_texte(spec)
        assert CARSO_DEFAUT.tableau.entete_fond in html
        assert "margin: 20.0mm 20.0mm 18.0mm 30.0mm" in html


class TestFacade:
    def test_formats_et_mime(self) -> None:
        # `xlsx` rejoint la liste à l'incrément 9 (ADR 0007) : un classeur est
        # désormais une sortie possible du même `DocumentSpec`.
        assert set(FORMATS_RENDUS) == {"docx", "pdf", "html", "xlsx"}
        assert extension_et_mime("DOCX") == (
            ".docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert extension_et_mime("xlsx")[0] == ".xlsx"
        with pytest.raises(ValidationError, match="inconnu"):
            rendre(_spec_complet(), format="odt")

    def test_html_est_un_document_autonome(self) -> None:
        contenu = rendre(_spec_complet(), format="html")
        assert contenu.startswith(b"<!DOCTYPE html>")
        assert b"charset=\"utf-8\"" in contenu


class TestDocxXml:
    """Couche XML isolée : ce que python-docx n'expose pas."""

    @staticmethod
    def _document() -> tuple[object, object]:
        import docx

        document = docx.Document()
        table = document.add_table(rows=2, cols=2)
        return document, table

    def test_largeurs_colonnes_et_mise_en_page_fixe(self) -> None:
        from docx.oxml.ns import qn

        _document_, table = self._document()
        mise_en_page_fixe(table)
        assert table._tbl.tblPr.find(qn("w:tblLayout")) is not None
        largeurs_colonnes(table, [60.0, 112.0])
        grille = table._tbl.find(qn("w:tblGrid"))
        assert [col.get(qn("w:w")) for col in grille.findall(qn("w:gridCol"))] == [
            "3402",
            "6350",
        ]
        with pytest.raises(ValidationError, match="largeurs"):
            largeurs_colonnes(table, [10.0])

    def test_trames_bordures_repetition(self) -> None:
        from docx.oxml.ns import qn

        _document_, table = self._document()
        ombrer_cellule(table.cell(0, 0), "#1F4E79")
        repetition_entete(table.rows[0])
        bordures_tableau(table, couleur="#B4C6E7", taille_pt=0.75)
        assert table.cell(0, 0)._tc.find(qn("w:tcPr")).find(qn("w:shd")) is not None
        assert table.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
        bordures = table._tbl.tblPr.find(qn("w:tblBorders"))
        assert len(bordures.findall(qn("w:top"))) == 1

    def test_espacements_filets_et_champ_page(self) -> None:
        import docx
        from docx.oxml.ns import qn

        document = docx.Document()
        paragraphe = document.add_paragraph("Texte")
        espacement_paragraphe(paragraphe, avant_pt=6, apres_pt=12, interligne=1.5)
        filet_haut_paragraphe(paragraphe, couleur="#BFBFBF")
        champ_numero_page(paragraphe)
        espacement = paragraphe._p.find(qn("w:pPr")).find(qn("w:spacing"))
        assert espacement.get(qn("w:before")) == "120"
        assert espacement.get(qn("w:line")) == "360"
        assert paragraphe._p.find(qn("w:pPr")).find(qn("w:pBdr")) is not None
        assert "PAGE" in paragraphe._p.xml


class TestStyleDeReference:
    """Le style est portable : un token peut venir d'un autre document."""

    def test_style_personnalise_applique(self) -> None:
        import docx
        from docx.shared import Pt

        style = CARSO_DEFAUT.model_copy(
            update={
                "titre": CARSO_DEFAUT.titre.model_copy(
                    update={"police": "Georgia", "taille_pt": 24.0, "couleur": "#0B2B4A"}
                ),
                "couverture": StyleCouverture(titre_taille_pt=30.0, couleur_accent="#0B2B4A"),
            }
        )
        spec = DocumentSpec(
            metadata=Metadonnees(titre="Titre stylé"),
            style=style,
            blocs=[Titre(niveau=1, texte="Section")],
        )
        document = docx.Document(io.BytesIO(rendre(spec, format="docx")))
        run = next(
            run
            for p in document.paragraphs
            if p.text == "Titre stylé"
            for run in p.runs
        )
        assert run.font.name == "Georgia"
        assert run.font.size == Pt(30.0)

    def test_generation_error_si_librairie_absente(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """Librairie absente ⇒ échec explicite, jamais un fichier vide."""
        import app.documents.renderers.docx as module

        def _echec() -> object:
            raise GenerationError("python-docx indisponible")

        monkeypatch.setattr(module, "_libs", _echec)
        with pytest.raises(GenerationError):
            rendre(_spec_complet(), format="docx", images=_images())
