"""Génération professionnelle DOCX/PDF — moteur vendu et tool CARSO.

Quatre points de contrôle :

1. **moteur** (``app.documents.professionnel``) : produit un DOCX lisible par
   ``python-docx`` et un PDF réellement paginé, refuse un contenu vide ou un
   vocabulaire hors bornes avec une erreur **explicite** ;
2. **tool** ``generer_document_professionnel`` : dépose en statut ``proposed``
   (règle 6 [C] : jamais ``approved`` depuis un agent), avec ancre métier réelle
   ou rangement ``non_classe``, et **ne renvoie aucun chemin disque** ;
3. **jumeau PDF** : un DOCX dépose aussi son PDF (même contenu, même pagination),
   pour que ``render_document`` — qui ne rend que des PDF — reste utilisable ;
4. **frontière de permission** : sans ``document_propose``, la génération est
   refusée ; et le module ``app.tools.documents`` n'importe pas ``playwright`` au
   chargement (le backend démarre sans l'extra ``documents_pro``).

La vérification qualité du moteur (avertissements) est reprise du module
d'origine dans ``test_documents_professionnel_verification.py``.

SQLite mémoire + stockage local temporaire (mêmes fixtures que
``test_tools_documents_html.py``).
"""

from __future__ import annotations

import io
import subprocess
import sys
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from uuid import UUID

import pytest
from docx import Document as LireDocx
from pypdf import PdfReader
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.application.services import DocumentService
from app.core.errors import ValidationError
from app.documents.professionnel import rendre_professionnel
from app.documents.professionnel.pdf_renderer import navigateur_disponible
from app.documents.storage import LocalDocumentStorage
from app.domain.base import Base
from app.domain.organization import Organisation
from app.tools.documents import build_document_tools
from app.tools.permissions import AgentIdentity, PermissionPolicy

IDENTITE = AgentIdentity(agent_id="agent_generateur_offre")
POLICY_DOC = PermissionPolicy(
    allowed_capabilities=frozenset({"document_read", "document_propose"})
)
POLICY_LECTURE = PermissionPolicy(allowed_capabilities=frozenset({"document_read"}))

#: Le navigateur est une dépendance d'exécution du PDF, pas des tests : la suite
#: reste verte sur une machine sans Edge/Chrome, et le test le DIT (jamais un
#: saut muet).
NAVIGATEUR_DISPONIBLE, RAISON_NAVIGATEUR = navigateur_disponible()
sans_navigateur = pytest.mark.skipif(
    not NAVIGATEUR_DISPONIBLE, reason=f"aucun navigateur Chromium : {RAISON_NAVIGATEUR}"
)

HTML_EXEMPLE = """
<section class="cover">
  <div class="cover-label">Rapport</div>
  <h1 class="cover-title">Bilan T3 2026</h1>
  <p class="cover-subtitle">Synthèse opérationnelle</p>
  <div class="cover-meta">Préparé pour <strong>Client SA</strong> · Octobre 2026</div>
</section>
<nav class="toc"><h2>Sommaire</h2></nav>
<div class="page-break"></div>
<h1>1. Synthèse</h1>
<p class="lead">Le trimestre dépasse l'objectif de 8 %.</p>
<table class="kpi"><tr>
  <td><span class="kpi-value">+24 %</span><span class="kpi-label">Croissance</span></td>
  <td><span class="kpi-value">1 250</span><span class="kpi-label">Clients</span></td>
</tr></table>
<div class="callout success"><strong>Recommandation :</strong> renforcer l'équipe PME.</div>
<h1>2. Détail</h1>
<table class="data">
  <thead><tr><th>Poste</th><th class="num">Montant</th></tr></thead>
  <tbody><tr><td>Formation</td><td class="num">4 800 €</td></tr></tbody>
  <tfoot><tr><td>Total</td><td class="num">4 800 €</td></tr></tfoot>
</table>
"""


@pytest.fixture
def session() -> Iterator[Session]:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as active:
        yield active
    engine.dispose()


@pytest.fixture
def storage(tmp_path: Path) -> LocalDocumentStorage:
    return LocalDocumentStorage(tmp_path / "carso", max_bytes=5 * 1024 * 1024)


@pytest.fixture
def service(session: Session, storage: LocalDocumentStorage) -> DocumentService:
    return DocumentService(session, storage=storage, actor_id="test")


@pytest.fixture
def organisation(session: Session) -> Organisation:
    organisation = Organisation(nom="Client Professionnel")
    session.add(organisation)
    session.flush()
    return organisation


def _tools(
    service: DocumentService,
    policy: PermissionPolicy = POLICY_DOC,
) -> dict[str, Any]:
    outils = build_document_tools(
        identite=IDENTITE,
        policy=policy,
        service_factory=lambda: service,
    )
    return {outil.name: outil for outil in outils}


def _payload(organisation: Organisation | None = None, **overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "contenu_html": HTML_EXEMPLE,
        "titre": "Bilan T3 2026",
        "format": "docx",
        "type_document": "rapport",
        "propose_par": "agent_generateur_offre",
    }
    if organisation is not None:
        base["ancre"] = "organisation"
        base["entity_id"] = str(organisation.id)
    base.update(overrides)
    return base


def _chaines(valeur: Any) -> Iterator[str]:
    """Toutes les chaînes d'une structure de retour (dict/list/str)."""
    if isinstance(valeur, str):
        yield valeur
    elif isinstance(valeur, dict):
        for element in valeur.values():
            yield from _chaines(element)
    elif isinstance(valeur, (list, tuple)):
        for element in valeur:
            yield from _chaines(element)


# --- Moteur -------------------------------------------------------------------


def test_moteur_produit_un_docx_lisible() -> None:
    rendu = rendre_professionnel(
        HTML_EXEMPLE, "Bilan T3 2026", format="docx", auteur="CARSO"
    )

    assert rendu.format == "docx"
    assert rendu.octets[:2] == b"PK"  # signature ZIP d'un .docx
    document = LireDocx(io.BytesIO(rendu.octets))
    textes = "\n".join(p.text for p in document.paragraphs)
    assert "1. Synthèse" in textes and "2. Détail" in textes
    # Les tableaux du thème sont de vrais tableaux Word, pas des paragraphes.
    assert document.tables
    assert rendu.avertissements == ()


def test_moteur_docx_n_a_pas_besoin_de_navigateur() -> None:
    """Le DOCX est du python-docx natif : sa production ne dépend pas de Chromium."""
    rendu = rendre_professionnel(
        "<h1>1. Note</h1><p>Corps.</p>", "Note", format="docx"
    )
    assert rendu.octets[:2] == b"PK"


@sans_navigateur
def test_moteur_produit_un_pdf_pagine() -> None:
    rendu = rendre_professionnel(HTML_EXEMPLE, "Bilan T3 2026", format="pdf")

    assert rendu.octets[:5] == b"%PDF-"
    assert rendu.jumeau_pdf is None  # un PDF n'a pas de jumeau
    assert rendu.pages is not None and rendu.pages >= 3  # couverture + 2 sections
    assert rendu.pages == len(PdfReader(io.BytesIO(rendu.octets)).pages)
    assert rendu.avertissements == ()


@sans_navigateur
def test_moteur_docx_rend_le_pdf_jumeau() -> None:
    rendu = rendre_professionnel(HTML_EXEMPLE, "Bilan T3 2026", format="docx")

    assert rendu.jumeau_pdf is not None
    assert rendu.jumeau_pdf[:5] == b"%PDF-"
    assert rendu.pages is not None
    assert rendu.pages == len(PdfReader(io.BytesIO(rendu.jumeau_pdf)).pages)


def test_moteur_refuse_un_contenu_vide() -> None:
    with pytest.raises(ValidationError):
        rendre_professionnel("   ", "Titre", format="pdf")


def test_moteur_refuse_un_format_hors_vocabulaire() -> None:
    with pytest.raises(ValidationError):
        rendre_professionnel("<h1>T</h1>", "Titre", format="xlsx")


def test_moteur_refuse_un_theme_inconnu() -> None:
    with pytest.raises(ValidationError):
        rendre_professionnel("<h1>T</h1>", "Titre", format="docx", theme="fluo")


def test_moteur_refuse_une_orientation_inconnue() -> None:
    with pytest.raises(ValidationError):
        rendre_professionnel("<h1>T</h1>", "Titre", format="docx", orientation="diagonal")


def test_moteur_retire_les_balises_interdites() -> None:
    """Le HTML de l'agent est nettoyé : le <style> fourni n'atteint pas le rendu."""
    rendu = rendre_professionnel(
        "<style>body{color:red}</style><script>alert(1)</script><h1>1. T</h1><p>Corps.</p>",
        "Titre",
        format="docx",
    )
    document = LireDocx(io.BytesIO(rendu.octets))
    texte = "\n".join(p.text for p in document.paragraphs)
    assert "alert(1)" not in texte and "color:red" not in texte
    assert "1. T" in texte


# --- Tool : dépôt HITL et permissions -----------------------------------------


def test_tool_depose_en_proposed_avec_ancre(
    service: DocumentService, organisation: Organisation
) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(organisation))

    assert sortie["statut"] == "proposed"  # HITL : jamais approved depuis un agent
    assert sortie["nom"].endswith(".docx")
    assert sortie["type_document"] == "rapport"
    document = service.obtenir(UUID(sortie["document_id"]))
    assert document.statut == "proposed"
    assert document.created_by == "agent_generateur_offre"
    assert document.organisation_id == organisation.id  # ancre réelle


def test_tool_sans_ancre_range_en_non_classe(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None))

    assert sortie.get("non_classe") is True
    assert "generated/" in sortie["message"]


def test_tool_refuse_sans_permission(
    service: DocumentService, organisation: Organisation
) -> None:
    outil = _tools(service, POLICY_LECTURE)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(organisation))

    assert sortie["ok"] is False
    assert sortie["error"]["code"] == "permission_denied"


def test_tool_ne_renvoie_aucun_chemin_disque(
    service: DocumentService, organisation: Organisation
) -> None:
    """Un chemin de fichier n'est jamais transmis à l'agent (AGENTS.md §7)."""
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(organisation))

    for valeur in _chaines(sortie):
        assert "\\" not in valeur, valeur
        assert "/tmp" not in valeur and "/carso" not in valeur, valeur


def test_tool_refuse_un_format_hors_vocabulaire(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None, format="xlsx"))

    assert sortie["ok"] is False


@sans_navigateur
def test_tool_pdf_depose_le_document(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None, format="pdf"))

    assert sortie["statut"] == "proposed"
    assert sortie["nom"].endswith(".pdf")
    assert sortie["pages"] >= 1
    assert sortie["documents"] == {"principal": sortie["document_id"], "format": "pdf"}


@sans_navigateur
def test_tool_docx_depose_aussi_le_pdf_jumeau(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None, format="docx"))

    assert sortie["documents"]["format"] == "docx"
    jumeau = sortie["documents"]["pdf"]
    assert jumeau != sortie["document_id"]
    document_jumeau = service.obtenir(UUID(jumeau))
    assert document_jumeau.nom.endswith(".pdf")
    assert document_jumeau.statut == "proposed"


def test_tool_remonte_les_avertissements(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(
        _payload(None, contenu_html='<style>p{color:red}</style><h1>1. Rapport</h1>')
    )

    assert any("<style>" in avertissement for avertissement in sortie["avertissements"])


def test_tool_derive_le_nom_du_titre_quand_il_est_omis(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None, titre="Compte rendu – Comité projet"))

    assert sortie["nom"] == "compte_rendu_comite_projet.docx"


# --- Démarrage sans l'extra ``documents_pro`` ---------------------------------


def test_le_backend_n_importe_pas_playwright_au_chargement() -> None:
    """``playwright`` est optionnel : son import est paresseux (dans le handler).

    Le backend doit démarrer sans l'extra ``documents_pro`` : si le module des
    tools documentaires importait le moteur professionnel au chargement, tous les
    autres tools documentaires deviendraient indisponibles sans ce paquet.
    """
    racine = Path(__file__).resolve().parents[1]
    code = (
        "import sys; import app.tools.documents; "
        "assert 'playwright' not in sys.modules, "
        "'playwright importe au chargement de app.tools.documents'"
    )
    resultat = subprocess.run(
        [sys.executable, "-c", code],
        cwd=racine,
        capture_output=True,
        text=True,
    )
    assert resultat.returncode == 0, resultat.stderr
