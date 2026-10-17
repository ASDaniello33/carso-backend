"""Présentation native PPTX (format ``pptx``) — modèle, rendu et tool.

Cinq points de contrôle :

1. **lecture** (``presentation.lire_presentation``) : le corps HTML de l'agent
   devient des diapositives typées (couverture, titre, lead, puces, KPI,
   tableau, figure, notes) — et refuse un HTML sans ``<section class="slide">`` ;
2. **rendu** (``rendre_professionnel(format="pptx")``) : le fichier produit est
   une vraie présentation PowerPoint relue par ``python-pptx`` — 16:9, une
   diapositive par section, notes et tableaux natifs — et passe la même
   validation d'extension qu'un dépôt humain (``inspect_file``) ;
3. **honnêteté de mise en page** : ce qui déborderait est REFUSÉ avec un message
   actionnable (jamais rogné en silence) ; un SVG sans ``resvg-py`` échoue
   explicitement (PowerPoint ne lit pas le SVG) ;
4. **tool** : le dépôt suit la règle HITL (statut ``proposed``, aucune ancre
   inventée, aucun chemin disque renvoyé) et annonce ses diapositives ;
5. **allowlist** : un fichier ``.pptx`` falsifié (contenu DOCX) est refusé.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from uuid import UUID

import pytest
from pptx import Presentation
from pptx.util import Emu
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.application.services import DocumentService
from app.core.errors import ValidationError
from app.documents.professionnel import (
    lire_presentation,
    rendre_professionnel,
    verifier_presentation,
)
from app.documents.storage import LocalDocumentStorage
from app.documents.validation import inspect_file
from app.domain.base import Base
from app.domain.organization import Organisation
from app.tools.documents import build_document_tools
from app.tools.permissions import AgentIdentity, PermissionPolicy

IDENTITE = AgentIdentity(agent_id="agent_assistant_formateur")
POLICY_DOC = PermissionPolicy(
    allowed_capabilities=frozenset({"document_read", "document_propose"})
)
POLICY_LECTURE = PermissionPolicy(allowed_capabilities=frozenset({"document_read"}))

DECK = """
<section class="cover">
  <div class="cover-label">Présentation</div>
  <h1 class="cover-title">Offre de formation DATA 2026</h1>
  <p class="cover-subtitle">Proposition technique</p>
  <div class="cover-meta">CARSO · Octobre 2026</div>
</section>
<section class="slide">
  <h1>Synthèse de la proposition</h1>
  <p class="lead">Un parcours de 5 jours aligné sur les lots sélectionnés.</p>
  <ul>
    <li>Première puce courte.</li>
    <li>Deuxième puce courte.</li>
    <li>Troisième puce courte.</li>
  </ul>
  <table class="kpi"><tr>
    <td><span class="kpi-value">5 j</span><span class="kpi-label">Durée</span></td>
    <td><span class="kpi-value">12</span><span class="kpi-label">Bénéficiaires</span></td>
  </tr></table>
  <aside class="notes">Rappeler l'ancrage sur le lot 1.</aside>
</section>
<section class="slide">
  <h1>Calendrier prévisionnel</h1>
  <table class="data">
    <thead><tr><th>Étape</th><th class="num">Jours</th></tr></thead>
    <tbody>
      <tr><td>Cadrage</td><td class="num">0,5</td></tr>
      <tr><td>Ateliers</td><td class="num">3</td></tr>
    </tbody>
  </table>
  <div class="callout warning">Échéance de remise : 30 novembre 2026.</div>
</section>
"""


# --- 1. Lecture du corps HTML --------------------------------------------------


def test_lire_presentation_interprete_le_deck() -> None:
    diapositives = lire_presentation(DECK)

    assert len(diapositives) == 3
    couverture, synthese, calendrier = diapositives
    assert couverture.est_titre
    assert couverture.titre == "Offre de formation DATA 2026"
    assert couverture.surtitre == "Présentation"
    assert couverture.sous_titre == "Proposition technique"
    assert couverture.meta.startswith("CARSO")
    assert synthese.titre == "Synthèse de la proposition"
    assert synthese.lead.startswith("Un parcours de 5 jours")
    assert syntheses_puces(synthese) == 3
    assert synthese.notes == "Rappeler l'ancrage sur le lot 1."
    assert any(bloc.genre == "kpi" and bloc.kpis for bloc in synthese.blocs)
    assert any(
        bloc.genre == "tableau"
        and bloc.tableau is not None
        and bloc.tableau.entetes == ("Étape", "Jours")
        for bloc in calendrier.blocs
    )
    assert any(bloc.genre == "encadre" and bloc.variante == "warning" for bloc in calendrier.blocs)


def syntheses_puces(diapositive: Any) -> int:
    return sum(len(bloc.items) for bloc in diapositive.blocs if bloc.genre == "puces")


def test_lire_presentation_refuse_un_html_sans_diapositive() -> None:
    with pytest.raises(ValidationError) as excinfo:
        lire_presentation("<h1>Un document ordinaire</h1><p>Pas un deck.</p>")
    assert "section" in str(excinfo.value.message)


def test_verifier_presentation_signale_les_diapositives_muettes() -> None:
    deck = """
    <section class="cover"><h1 class="cover-title">Titre</h1></section>
    <section class="slide"><p>Un paragraphe sans titre ni notes.</p></section>
    """
    constats = verifier_presentation(lire_presentation(deck))
    assert any("aucun titre" in c for c in constats)
    assert any("notes présentateur" in c for c in constats)
    assert not any("aucun contenu" in c for c in constats)  # le paragraphe compte


# --- 2. Rendu PPTX réel --------------------------------------------------------


def _ouvrir(octets: bytes) -> Presentation:
    return Presentation(io.BytesIO(octets))


def test_pptx_produit_une_presentation_lisible() -> None:
    rendu = rendre_professionnel(DECK, "Offre de formation DATA 2026", format="pptx")

    assert rendu.format == "pptx"
    assert rendu.octets[:2] == b"PK"  # signature ZIP d'un OOXML
    assert rendu.pages == 3  # une « page » = une diapositive
    assert rendu.jumeau_pdf is None
    presentation = _ouvrir(rendu.octets)
    assert presentation.slide_width == Emu(12_192_000)  # 16:9 canonique
    assert presentation.slide_height == Emu(6_858_000)
    assert len(presentation.slides) == 3
    # La couverture porte le titre ; la diapositive 2, ses puces et ses notes.
    textes_couverture = "\n".join(
        forme.text_frame.text
        for forme in presentation.slides[0].shapes
        if forme.has_text_frame
    )
    assert "Offre de formation DATA 2026" in textes_couverture
    diapo2 = presentation.slides[1]
    assert diapo2.has_notes_slide
    assert "lot 1" in diapo2.notes_slide.notes_text_frame.text
    # Un table.data devient un vrai tableau PowerPoint, pas des images.
    assert any(forme.has_table for forme in presentation.slides[2].shapes)


def test_pptx_passe_la_validation_de_depôt(tmp_path: Path) -> None:
    rendu = rendre_professionnel(DECK, "Deck", format="pptx")
    chemin = tmp_path / "deck.pptx"
    chemin.write_bytes(rendu.octets)

    regle = inspect_file(chemin, nom_fichier="deck.pptx")
    assert regle.mime_type == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )


def test_pptx_ne_fabrique_pas_un_xlsx(tmp_path: Path) -> None:
    """Un deck dépôt ses tableaux en formes PowerPoint : l'extension reste .pptx."""
    rendu = rendre_professionnel(DECK, "Deck", format="pptx", auteur="CARSO")
    presentation = _ouvrir(rendu.octets)
    assert presentation.core_properties.author == "CARSO"
    with zipfile.ZipFile(io.BytesIO(rendu.octets)) as archive:
        assert "ppt/presentation.xml" in archive.namelist()
        assert "xl/workbook.xml" not in archive.namelist()


def test_pptx_refuse_un_contenu_vide_ou_un_theme_inconnu() -> None:
    with pytest.raises(ValidationError):
        rendre_professionnel("   ", "Titre", format="pptx")
    with pytest.raises(ValidationError):
        rendre_professionnel(DECK, "Titre", format="pptx", theme="fluo")


# --- 3. Honnêteté de mise en page ---------------------------------------------


def test_pptx_refuse_une_diapositive_surchargee() -> None:
    surchargee = f"""
    <section class="slide"><h1>Trop de texte</h1>
      <ul><li>{"Une puce qui ne finit jamais. " * 200}</li></ul>
    </section>
    """
    with pytest.raises(ValidationError) as excinfo:
        rendre_professionnel(surchargee, "Deck", format="pptx")
    assert "trop chargée" in str(excinfo.value.message)


def test_pptx_refuse_un_tableau_trop_haut() -> None:
    lignes = "".join(f"<tr><td>Ligne {i}</td><td class='num'>{i}</td></tr>" for i in range(20))
    deck = (
        '<section class="slide"><h1>Tableau énorme</h1>'
        f'<table class="data"><thead><tr><th>Poste</th><th>N</th></tr></thead>'
        f"<tbody>{lignes}</tbody></table></section>"
    )
    with pytest.raises(ValidationError) as excinfo:
        rendre_professionnel(deck, "Deck", format="pptx")
    assert "trop haut" in str(excinfo.value.message)


def test_pptx_refuse_une_image_absolue_manquante() -> None:
    deck = (
        '<section class="slide"><h1>Figure manquante</h1>'
        '<figure><img src="inexistant.png"><figcaption>Légende</figcaption></figure>'
        "</section>"
    )
    with pytest.raises(ValidationError) as excinfo:
        rendre_professionnel(deck, "Deck", format="pptx")
    assert "Image introuvable" in str(excinfo.value.message)


def test_pptx_refuse_une_image_distante() -> None:
    deck = (
        '<section class="slide"><h1>Image distante</h1>'
        '<figure><img src="https://exemple.fr/logo.png"></figure></section>'
    )
    with pytest.raises(ValidationError) as excinfo:
        rendre_professionnel(deck, "Deck", format="pptx")
    assert "distante" in str(excinfo.value.message)


def test_pptx_sans_resvg_echoue_explicitement(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """PowerPoint ne lit pas le SVG : sans rasteriseur, on refuse — jamais un trou."""
    import app.documents.professionnel.rendu as rendu_module

    monkeypatch.setattr(rendu_module, "rasterisation_disponible", lambda: False)
    deck = (
        '<section class="slide"><h1>Schéma</h1>'
        '<figure><svg viewBox="0 0 100 50" width="200">'
        '<rect width="80" height="30" fill="#2563EB"/></svg></figure>'
        "</section>"
    )
    with pytest.raises(ValidationError) as excinfo:
        rendre_professionnel(deck, "Deck", format="pptx")
    assert "resvg-py" in str(excinfo.value.message)


def test_pptx_integre_un_graphique_en_image() -> None:
    """Une ancre data-graphique devient une image PNG (relue comme Picture)."""
    deck = (
        '<section class="slide"><h1>Chiffres clés</h1>'
        "<figure data-graphique='repartition'></figure></section>"
    )
    rendu = rendre_professionnel(
        deck,
        "Deck",
        format="pptx",
        graphiques={
            "repartition": {
                "type": "bar",
                "labels": ["T1", "T2", "T3"],
                "series": {"Formations": [3, 5, 8]},
                "width": 640,
                "height": 400,
            }
        },
    )
    presentation = _ouvrir(rendu.octets)
    images = [
        forme
        for forme in presentation.slides[0].shapes
        if forme.shape_type is not None and "PICTURE" in str(forme.shape_type)
    ]
    assert images, "le graphique devrait être une image PowerPoint"


# --- Placement des figures (largeur réelle de la diapositive) ------------------


_SVG_BANDEAU = (
    '<figure><svg viewBox="0 0 800 230" width="800" height="230" '
    'xmlns="http://www.w3.org/2000/svg">'
    '<rect x="10" y="40" width="140" height="90" fill="#16A34A"/>'
    '<rect x="330" y="40" width="140" height="90" fill="#22C55E"/>'
    '<rect x="650" y="40" width="140" height="90" fill="#16A34A"/>'
    "</svg><figcaption>Cycle en 5 phases</figcaption></figure>"
)


def _largeur_picture(octets: bytes) -> float:
    presentation = _ouvrir(octets)
    for slide in presentation.slides:
        for forme in slide.shapes:
            if forme.shape_type is not None and "PICTURE" in str(forme.shape_type):
                return forme.width / 914400
    pytest.fail("aucune image dans le deck")


def test_pptx_figure_seule_pleine_largeur() -> None:
    """Une figure sans texte occupe toute la largeur utile, pas une demi-colonne."""
    deck = f'<section class="slide"><h1>Cycle</h1>{_SVG_BANDEAU}</section>'
    rendu = rendre_professionnel(deck, "Deck", format="pptx")

    assert _largeur_picture(rendu.octets) == pytest.approx(11.833, abs=0.05)


def test_pptx_bandeau_large_empile_sous_le_texte() -> None:
    """Une frise (ratio h/l < 0,45) ne s'écrase pas dans la colonne de droite."""
    deck = (
        '<section class="slide"><h1>Cycle</h1>'
        f"<p>Une phrase de contexte.</p>{_SVG_BANDEAU}</section>"
    )
    rendu = rendre_professionnel(deck, "Deck", format="pptx")

    assert _largeur_picture(rendu.octets) == pytest.approx(11.833, abs=0.05)


def test_pptx_figure_haute_restee_en_colonne() -> None:
    """Figure carrée avec du texte : split 50/50 (une frise non, un schéma oui)."""
    carre = _SVG_BANDEAU.replace(
        'viewBox="0 0 800 230" width="800" height="230"',
        'viewBox="0 0 400 300" width="400" height="300"',
    )
    deck = (
        '<section class="slide"><h1>Cycle</h1>'
        f"<p>Une phrase de contexte.</p>{carre}</section>"
    )
    rendu = rendre_professionnel(deck, "Deck", format="pptx")

    assert _largeur_picture(rendu.octets) == pytest.approx(6.13, abs=0.3)


# --- 4. Tool : dépôt HITL ------------------------------------------------------


@pytest.fixture
def session() -> Iterator[Session]:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as active:
        yield active
    engine.dispose()


@pytest.fixture
def service(session: Session, tmp_path: Path) -> DocumentService:
    stockage = LocalDocumentStorage(tmp_path / "carso", max_bytes=10 * 1024 * 1024)
    return DocumentService(session, storage=stockage, actor_id="test")


def _tools(service: DocumentService, policy: PermissionPolicy = POLICY_DOC) -> dict[str, Any]:
    outils = build_document_tools(
        identite=IDENTITE,
        policy=policy,
        service_factory=lambda: service,
    )
    return {outil.name: outil for outil in outils}


def _payload(organisation: Organisation | None = None, **overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "contenu_html": DECK,
        "titre": "Offre de formation DATA 2026",
        "format": "pptx",
        "type_document": "rapport",
        "propose_par": "agent_assistant_formateur",
    }
    if organisation is not None:
        base["ancre"] = "organisation"
        base["entity_id"] = str(organisation.id)
    base.update(overrides)
    return base


@pytest.fixture
def organisation(session: Session) -> Organisation:
    organisation = Organisation(nom="Client Professionnel")
    session.add(organisation)
    session.flush()
    return organisation


def test_tool_pptx_depose_en_proposed_avec_ancre(
    service: DocumentService, organisation: Organisation
) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(organisation))

    assert sortie["statut"] == "proposed"  # HITL : jamais approved depuis un agent
    assert sortie["nom"].endswith(".pptx")
    assert sortie["documents"]["format"] == "pptx"
    assert sortie["diapositives"] == 3
    assert "limite_apercu" in sortie  # render_document ne lit que les PDF
    document = service.obtenir(UUID(sortie["document_id"]))
    assert document.organisation_id == organisation.id
    assert document.mime_type.endswith("presentationml.presentation")


def test_tool_pptx_sans_ancre_range_en_non_classe(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None))

    assert sortie.get("non_classe") is True
    assert "generated/" in sortie["message"]


def test_tool_pptx_refuse_sans_permission(service: DocumentService) -> None:
    outil = _tools(service, POLICY_LECTURE)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None))

    assert sortie["ok"] is False
    assert sortie["error"]["code"] == "permission_denied"


def test_tool_pptx_ne_renvoie_aucun_chemin_disque(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None))

    assert "chemin" not in sortie
    assert "C:\\" not in sortie["message"]
    document = service.obtenir(UUID(sortie["document_id"]))
    assert document.nom.endswith(".pptx")


def test_tool_pptx_refuse_une_diapositive_surchargee(service: DocumentService) -> None:
    surchargee = (
        '<section class="slide"><h1>Trop de texte</h1>'
        f"<ul><li>{'Texte interminable. ' * 200}</li></ul></section>"
    )
    outil = _tools(service)["generer_document_professionnel"]
    sortie = outil.invoke(_payload(None, contenu_html=surchargee))

    assert sortie["ok"] is False
    assert "trop chargée" in str(sortie["error"]["message"])


# --- 5. Allowlist : un .pptx falsifié est refusé -------------------------------


def test_inspect_file_refuse_un_docx_renomme_pptx(tmp_path: Path) -> None:
    falsifie = tmp_path / "faux.pptx"
    with zipfile.ZipFile(falsifie, "w") as archive:
        archive.writestr("word/document.xml", "<w:document/>")
        archive.writestr("[Content_Types].xml", "<Types/>")

    with pytest.raises(ValidationError):
        inspect_file(falsifie, nom_fichier="faux.pptx")
