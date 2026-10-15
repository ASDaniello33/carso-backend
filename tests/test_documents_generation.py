"""Lot L0 — génération documentaire DOCX/XLSX/PDF (octets purs, sans stockage).

Le pipeline [D2] impose que tout document d'agent passe par DocumentService ;
ici on vérifie uniquement les moteurs : contenu réel lisible par les mêmes
librairies que ``document_fixtures`` (pypdf, python-docx, openpyxl).
"""

import io
from decimal import Decimal

import pytest

from app.core.errors import GenerationError
from app.documents.generation import GenerateurDocument, build_docx, build_pdf, build_xlsx


class TestBuildDocx:
    def test_contenu_lisible(self) -> None:
        import docx

        octets = build_docx(
            titre="Fiche technique",
            paragraphes=["Mission M-2026-9", "Formateur : Sara Alaoui"],
            tableau=[["Nom", "Rôle"], ["Alaoui", "formateur"]],
        )

        document = docx.Document(io.BytesIO(octets))
        textes = [p.text for p in document.paragraphs]
        assert any("Fiche technique" in t for t in textes)
        assert any("Sara Alaoui" in t for t in textes)
        assert document.tables[0].cell(1, 0).text == "Alaoui"

    def test_sans_tableau(self) -> None:
        octets = build_docx(titre="Checklist", paragraphes=["item 1"])
        assert octets[:2] == b"PK"  # DOCX = archive ZIP


class TestBuildXlsx:
    def test_liste_beneficiaires(self) -> None:
        from openpyxl import load_workbook

        octets = build_xlsx(
            titre_feuille="Bénéficiaires",
            lignes=[["Nom", "Email"], ["Karimi", "karimi@ex.ma"]],
        )

        classeur = load_workbook(io.BytesIO(octets))
        onglet = classeur.active
        assert onglet.title == "Bénéficiaires"
        assert onglet.cell(1, 1).value == "Nom"
        assert onglet.cell(2, 1).value == "Karimi"


class TestBuildPdf:
    def test_texte_extractible(self) -> None:
        from pypdf import PdfReader

        octets = build_pdf(
            titre="Fiche de présence",
            paragraphes=["Session confirmée"],
            tableau=[["Bénéficiaire", "Signature"], ["Karimi", ""]],
        )

        lecteur = PdfReader(io.BytesIO(octets))
        texte = lecteur.pages[0].extract_text()
        assert "Fiche de présence" in texte
        assert "Karimi" in texte

    def test_titre_vide_refuse(self) -> None:
        with pytest.raises(GenerationError, match="titre"):
            build_pdf(titre="   ")

    def test_montant_decimal_affiche(self) -> None:
        from pypdf import PdfReader

        octets = build_pdf(
            titre="Budget",
            tableau=[["Ligne", "Montant"], ["Salle", Decimal("1500.50")]],
        )
        texte = PdfReader(io.BytesIO(octets)).pages[0].extract_text()
        assert "1500.50" in texte


class TestGenerateurDocument:
    def test_facade_trois_formats(self) -> None:
        generateur = GenerateurDocument()
        assert build_docx(titre="A")  # sanity
        assert generateur.docx(titre="A")[:2] == b"PK"
        assert generateur.xlsx(titre_feuille="S", lignes=[["h"]])
        assert generateur.pdf(titre="A").startswith(b"%PDF")
