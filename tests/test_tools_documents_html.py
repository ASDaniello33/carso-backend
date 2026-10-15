"""Tools documentaires HTML — ``generer_document_html`` et moteur de conversion.

Vérifie trois choses :

1. le moteur ``app.documents.html_to_docx`` produit un DOCX valide depuis du
   HTML+CSS (et refuse un HTML vide avec une erreur explicite) ;
2. le tool ``generer_document_html`` dépose en statut ``proposed`` (règle 6
   [C] : jamais ``approved`` depuis un agent), avec ancre métier réelle ou
   rangement ``non_classe`` ;
3. la frontière de permission : sans ``document_propose``, la génération est
   refusée ; le format ``pdf`` passe par la structure extraite du HTML.

SQLite mémoire + storage local temporaire (mêmes fixtures que
``test_tools_documents.py``).
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from uuid import UUID

import pytest
from docx import Document as LireDocx
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.application.services import DocumentService
from app.documents.html_to_docx import ConversionHtmlError, html_string_to_docx_bytes
from app.documents.storage import LocalDocumentStorage
from app.domain.base import Base
from app.domain.organization import Organisation
from app.tools.documents import build_document_tools
from app.tools.permissions import AgentIdentity, PermissionPolicy

IDENTITE = AgentIdentity(agent_id="agent_generaliste")
POLICY_DOC = PermissionPolicy(
    allowed_capabilities=frozenset({"document_read", "document_propose"})
)
POLICY_LECTURE = PermissionPolicy(allowed_capabilities=frozenset({"document_read"}))

HTML_EXEMPLE = """
<html><body>
<h1 style="color:#1a5276;">Note de synthèse</h1>
<p style="text-align:justify;">Un <b>paragraphe</b> avec de l'<i>italique</i>.</p>
<div style="background-color:#eaf2f8; border:1px solid #154360; padding:8px;">
  <p>Encadré important.</p>
</div>
<table>
  <tr><th style="background-color:#d6eaf8;">Poste</th><th>Montant</th></tr>
  <tr><td>Formation</td><td>150 000</td></tr>
</table>
<ul><li>Premier point</li><li>Second point</li></ul>
</body></html>
"""


@pytest.fixture
def session() -> Iterator[Session]:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as active:
        yield active
    engine.dispose()


@pytest.fixture
def storage(tmp_path: Path) -> LocalDocumentStorage:
    return LocalDocumentStorage(tmp_path / "carso", max_bytes=2 * 1024 * 1024)


@pytest.fixture
def service(session: Session, storage: LocalDocumentStorage) -> DocumentService:
    return DocumentService(session, storage=storage, actor_id="test")


@pytest.fixture
def organisation(session: Session) -> Organisation:
    organisation = Organisation(nom="Client HTML")
    session.add(organisation)
    session.flush()
    return organisation


def _tools(
    service: DocumentService,
    policy: PermissionPolicy = POLICY_DOC,
) -> dict[str, Any]:
    outils = build_document_tools(
        identite=IDENTITE,
        policy=policy,
        service_factory=lambda: service,
    )
    return {outil.name: outil for outil in outils}


def _payload(organisation: Organisation | None = None, **overrides: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "format": "docx",
        "html": HTML_EXEMPLE,
        "nom_fichier": "note_synthese",
        "type_document": "rapport",
        "propose_par": "agent_generaliste",
    }
    if organisation is not None:
        base["ancre"] = "organisation"
        base["entity_id"] = str(organisation.id)
    base.update(overrides)
    return base


# --- Moteur de conversion -----------------------------------------------------


def test_moteur_produit_un_docx_valide() -> None:
    octets = html_string_to_docx_bytes(HTML_EXEMPLE)
    assert octets[:2] == b"PK"  # signature ZIP d'un .docx
    document = LireDocx(io.BytesIO(octets))
    textes = "\n".join(p.text for p in document.paragraphs)
    assert "Note de synthèse" in textes
    assert "Premier point" in textes  # liste rendue
    # Le <table> et l'encadré (div à fond + bordure) sont rendus en tableaux.
    textes_tables = "\n".join(
        cellule.text
        for table in document.tables
        for ligne in table.rows
        for cellule in ligne.cells
    )
    assert "Encadré important." in textes_tables
    assert "Formation" in textes_tables and "150 000" in textes_tables


def test_moteur_refuse_html_vide() -> None:
    with pytest.raises(ConversionHtmlError):
        html_string_to_docx_bytes("   \n  ")


def test_moteur_refuse_html_sans_contenu() -> None:
    with pytest.raises(ConversionHtmlError):
        html_string_to_docx_bytes("<html><body></body></html>")


# --- Tool : dépôt HITL et permissions -----------------------------------------


def test_generation_html_en_statut_proposed_avec_ancre(
    service: DocumentService, organisation: Organisation
) -> None:
    outil = _tools(service)["generer_document_html"]
    sortie = outil.invoke(_payload(organisation))

    assert sortie["statut"] == "proposed"  # HITL : jamais approved depuis un agent
    assert sortie["nom"] == "note_synthese.docx"
    document = service.obtenir(UUID(sortie["document_id"]))
    assert document.statut == "proposed"
    assert document.created_by == "agent_generaliste"
    assert document.organisation_id == organisation.id  # ancre réelle


def test_generation_html_sans_ancre_est_non_classe(
    service: DocumentService,
) -> None:
    outil = _tools(service)["generer_document_html"]
    sortie = outil.invoke(_payload(None))

    assert sortie.get("non_classe") is True
    assert "generated/" in sortie["message"]


def test_generation_html_refusee_sans_permission(
    service: DocumentService, organisation: Organisation
) -> None:
    outil = _tools(service, POLICY_LECTURE)["generer_document_html"]
    sortie = outil.invoke(_payload(organisation))
    assert sortie["ok"] is False
    assert sortie["error"]["code"] == "permission_denied"


def test_generation_html_pdf(service: DocumentService) -> None:
    outil = _tools(service)["generer_document_html"]
    sortie = outil.invoke(_payload(None, format="pdf"))

    assert sortie["statut"] == "proposed"
    assert sortie["nom"] == "note_synthese.pdf"
    assert "note_pdf" in sortie  # limite honnête documentée dans la réponse
    document = service.obtenir(UUID(sortie["document_id"]))
    assert document.statut == "proposed"


def test_generation_html_format_inconnu_refuse(
    service: DocumentService,
) -> None:
    outil = _tools(service)["generer_document_html"]
    sortie = outil.invoke(_payload(None, format="xlsx"))
    assert sortie["ok"] is False


# --- Fusions HTML (colspan/rowspan) : l'ancien échec tool call ---------------


HTML_FUSIONS = """
<html><body>
<h1>Budget</h1>
<table>
  <tr><th>Poste</th><th>Quantité</th><th>Prix unitaire</th><th>Total</th></tr>
  <tr><td>Formation</td><td>5</td><td>100</td><td>500</td></tr>
  <tr><td colspan="4">Total général</td></tr>
</table>
</body></html>
"""


def test_pdf_avec_ligne_de_total_colspan_reussit(
    service: DocumentService,
) -> None:
    """Le bug « Tableau, ligne 3 : 1 colonne(s) non couverte(s) (colonnes [4]) ».

    La ligne de total ``<td colspan="4">`` produisait une grille de 1 cellule
    dans un tableau à 4 colonnes : le validateur du DocumentSpec refusait et
    l'appel tool échouait. La grille extraite couvre désormais toute la
    largeur — le PDF se génère.
    """
    outil = _tools(service)["generer_document_html"]
    sortie = outil.invoke(_payload(None, html=HTML_FUSIONS, format="pdf"))
    assert sortie["statut"] == "proposed", sortie.get("message")
    assert sortie["nom"].endswith(".pdf")


def test_docx_colspan_produit_une_fusion_reelle(
    service: DocumentService,
) -> None:
    """Le DOCX fusionne réellement les colonnes (avant : décalage silencieux)."""
    from docx import Document as LireDocxLocal

    octets = html_string_to_docx_bytes(HTML_FUSIONS)
    document = LireDocxLocal(io.BytesIO(octets))
    table = document.tables[0]
    # 3 lignes, 4 colonnes de grille.
    assert len(table.rows) == 3
    assert len(table.columns) == 4
    # La cellule de total fusionne les 4 colonnes : les 4 accès renvoient le
    # même contenu (comportement python-docx d'une fusion).
    total = table.cell(2, 0).text
    assert "Total général" in total
    assert table.cell(2, 3).text == total


def test_docx_rowspan_produit_une_fusion_verticale() -> None:
    """Un ``rowspan="2"`` fusionne les deux lignes (fidélité au HTML)."""
    from docx import Document as LireDocxLocal

    html = (
        "<html><body><table>"
        "<tr><th>Zone</th><th>Détail</th></tr>"
        "<tr><td rowspan=\"2\">Analamanga</td><td>Site 1</td></tr>"
        "<tr><td>Site 2</td></tr>"
        "</table></body></html>"
    )
    document = LireDocxLocal(io.BytesIO(html_string_to_docx_bytes(html)))
    table = document.tables[0]
    # La cellule « Analamanga » occupe (1,0) et (2,0) : même objet fusionné.
    assert "Analamanga" in table.cell(1, 0).text
    assert table.cell(2, 0).text == table.cell(1, 0).text
    assert table.cell(2, 1).text == "Site 2"


def test_pdf_rowspan_ne_provoque_pas_de_colonne_non_couverte(
    service: DocumentService,
) -> None:
    """Le chemin PDF aplatit les rowspan sans casser la grille."""
    outil = _tools(service)["generer_document_html"]
    html = (
        "<html><body><table>"
        "<tr><th>Zone</th><th>Détail</th></tr>"
        "<tr><td rowspan=\"2\">Analamanga</td><td>Site 1</td></tr>"
        "<tr><td>Site 2</td></tr>"
        "</table></body></html>"
    )
    sortie = outil.invoke(_payload(None, html=html, format="pdf"))
    assert sortie["statut"] == "proposed", sortie.get("message")


HTML_FUSIONS_CROISEES = """
<html><body><table>
  <tr><th colspan="2">En-tête à cheval</th><th>C</th></tr>
  <tr><td rowspan="2">Zone</td><td>a1</td><td>c1</td></tr>
  <tr><td>b2</td><td>c2</td></tr>
</table></body></html>
"""


def test_docx_fusions_croisees_gardent_le_contenu_de_chaque_cellule() -> None:
    """Horizontale et verticale dans le même tableau : rien n'est décalé.

    La grille est désormais lue **après** les fusions (une seule dérivation) :
    ce test verrouille l'indexation — une cellule fusionnée est répétée sur les
    positions qu'elle couvre, les autres gardent leur propre contenu.
    """
    document = LireDocx(io.BytesIO(html_string_to_docx_bytes(HTML_FUSIONS_CROISEES)))
    table = document.tables[0]
    assert (len(table.rows), len(table.columns)) == (3, 3)
    assert table.cell(0, 0).text == "En-tête à cheval"
    assert table.cell(0, 1).text == table.cell(0, 0).text
    assert table.cell(0, 2).text == "C"
    assert table.cell(1, 0).text == "Zone"
    assert table.cell(2, 0).text == table.cell(1, 0).text
    assert [table.cell(i, j).text for i, j in ((1, 1), (1, 2), (2, 1), (2, 2))] == [
        "a1",
        "c1",
        "b2",
        "c2",
    ]


def _html_grand_tableau(nb_lignes: int = 60, nb_colonnes: int = 6) -> str:
    """Un tableau réaliste : en-tête à cheval + ligne de total (fusions)."""
    entete = "<th colspan=\"2\">Poste</th>" + "".join(
        f"<th>Colonne {c}</th>" for c in range(2, nb_colonnes)
    )
    lignes = "".join(
        "<tr>"
        + "".join(f"<td>Cellule {i}-{c}</td>" for c in range(nb_colonnes))
        + "</tr>"
        for i in range(nb_lignes)
    )
    total = f'<tr><td colspan="{nb_colonnes}">Total général</td></tr>'
    return (
        f"<html><body><h1>Rapport</h1><table><thead><tr>{entete}</tr></thead>"
        f"<tbody>{lignes}{total}</tbody></table></body></html>"
    )


def test_la_grille_du_tableau_n_est_pas_derivee_une_fois_par_cellule(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Non-régression du gel d'interface du 07/10 (HTML → DOCX).

    ``Table.cell(i, j)`` de python-docx redérive la grille entière **à chaque
    appel** : une boucle cellule par cellule coûtait donc O(nb_cellules²), et un
    rapport de tableaux de 30 lignes figeait la requête 16 s. Le rendu doit
    dériver la grille un nombre constant de fois (une par fusion, plus une pour
    le contenu) — jamais une par cellule. Le compteur est un garde-fou
    déterministe, à la place d'un test de durée qui serait instable en CI.
    """
    import docx.table as module_table

    derivee = module_table.Table._cells
    appels = {"nb": 0}

    class _GrilleCountee:
        def __get__(self, instance: Any, owner: type | None = None) -> Any:
            if instance is None:
                return self
            appels["nb"] += 1
            return derivee.__get__(instance, owner)

    monkeypatch.setattr(module_table.Table, "_cells", _GrilleCountee())

    nb_cellules = 62 * 6
    octets = html_string_to_docx_bytes(_html_grand_tableau())
    derivations = appels["nb"]  # mesuré pendant le rendu, avant relecture du DOCX
    document = LireDocx(io.BytesIO(octets))

    assert derivations < nb_cellules // 10, (
        f"la grille du tableau est dérivée {derivations} fois pour {nb_cellules} "
        "cellules : le rendu est redevenu quadratique (gel de l'interface)"
    )

    table = document.tables[0]
    assert len(table.rows) == 62
    assert table.cell(1, 5).text == "Cellule 0-5"
    assert table.cell(60, 0).text == "Cellule 59-0"
    assert table.cell(61, 5).text == "Total général"
    assert table.cell(0, 0).text == table.cell(0, 1).text == "Poste"


# --- type_document : alias de profils et erreur pédagogique -------------------


def test_type_document_alias_profil_est_converti(
    service: DocumentService,
) -> None:
    """'offre_financiere' (profil) → type 'offre' : plus d'échec tool call."""
    outil = _tools(service)["generer_document_html"]
    sortie = outil.invoke(
        _payload(None, type_document="offre_financiere")
    )
    assert sortie["statut"] == "proposed"
    assert sortie["type_document"] == "offre"
    # Le retour signale la conversion : le modèle s'aligne au coup suivant.
    assert "type_document_convertis" in sortie
    assert "offre" in sortie["type_document_convertis"]


def test_type_document_inconnu_renvoie_une_erreur_pedagogique(
    service: DocumentService,
) -> None:
    """L'erreur liste les valeurs valides et explique la confusion profil/type."""
    outil = _tools(service)["generer_document_html"]
    sortie = outil.invoke(
        _payload(None, type_document="rapport_financier")
    )
    assert sortie["ok"] is False
    message = sortie["error"]["message"]
    assert "Type de document inconnu" in message
    details = sortie["error"].get("details", {})
    assert "offre" in details.get("types_valides", [])
    assert "profil" in details.get("suggestion", "")
