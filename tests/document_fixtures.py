"""Générateurs de vrais fichiers documentaires pour les tests.

Les tests documentaires ne s'appuient pas sur des octets factices : les formats
sont produits par les mêmes bibliothèques que celles utilisées en production
(pypdf, python-docx, openpyxl), ce qui vérifie réellement le couple
validation ↔ extraction. Les images PNG sont fabriquées avec ``zlib`` et
``struct`` (stdlib) : aucun besoin de Pillow pour les tests.

Ce module n'est pas un module de tests (nom sans préfixe ``test_``).
"""

from __future__ import annotations

import io
import struct
import zlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

_CONTENU_PDF = "BT /F1 12 Tf 20 200 Td ({texte}) Tj ET"


def build_pdf(texte: str = "MISSION CARSO 2026") -> bytes:
    """PDF d'une page contenant ``texte`` (extractible par pypdf)."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)

    contenu = DecodedStreamObject()
    contenu.set_data(_CONTENU_PDF.format(texte=texte).encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(contenu)

    police = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(police)})}
    )

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_pdf_sans_texte() -> bytes:
    """PDF d'une page blanche : aucun texte extractible (cas « scanné »)."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx(*, paragraphes: list[str], tableau: list[list[str]] | None = None) -> bytes:
    """DOCX avec paragraphes et, optionnellement, un tableau."""
    import docx

    document = docx.Document()
    for paragraphe in paragraphes:
        document.add_paragraph(paragraphe)
    if tableau:
        table = document.add_table(rows=len(tableau), cols=len(tableau[0]))
        for index_ligne, ligne in enumerate(tableau):
            for index_colonne, valeur in enumerate(ligne):
                table.cell(index_ligne, index_colonne).text = valeur

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_xlsx(*, feuille: str, lignes: list[list[object]]) -> bytes:
    """Classeur XLSX d'une feuille (en-tête inclus : liste de bénéficiaires type)."""
    classeur = Workbook()
    onglet = classeur.active
    onglet.title = feuille
    for ligne in lignes:
        onglet.append(ligne)

    buffer = io.BytesIO()
    classeur.save(buffer)
    return buffer.getvalue()


def build_png(
    largeur: int = 4,
    hauteur: int = 3,
    couleur: tuple[int, int, int] = (200, 30, 30),
) -> bytes:
    """PNG truecolor minimaliste, valide (en-tête + IHDR + IDAT + IEND)."""

    def chunk(tag: bytes, donnees: bytes) -> bytes:
        return (
            struct.pack(">I", len(donnees))
            + tag
            + donnees
            + struct.pack(">I", zlib.crc32(tag + donnees) & 0xFFFFFFFF)
        )

    brut = b"".join(b"\x00" + bytes(couleur) * largeur for _ in range(hauteur))
    return b"".join(
        [
            b"\x89PNG\r\n\x1a\n",
            chunk(b"IHDR", struct.pack(">IIBBBBB", largeur, hauteur, 8, 2, 0, 0, 0)),
            chunk(b"IDAT", zlib.compress(brut)),
            chunk(b"IEND", b""),
        ]
    )


def build_docx_riche() -> bytes:
    """DOCX réaliste : titres hiérarchisés, styles, tableau, image, en-tête, saut.

    Sert aux tests de **lecture structurée** et d'**extraction de style** :
    hiérarchie H1 > H2 > paragraphe > tableau, largeurs de colonnes, image,
    en-tête/pied et saut de page explicite.
    """
    import docx
    from docx.shared import Mm, Pt

    document = docx.Document()
    document.add_heading("Offre technique — Lot 3", level=1)
    document.add_paragraph("Contexte de la mission et attentes du commanditaire.")
    document.add_heading("Méthodologie", level=2)
    document.add_paragraph("Une démarche en trois phases, animée par deux formateurs.")
    document.add_paragraph("Phase 1 — cadrage", style="List Bullet")
    document.add_paragraph("Phase 2 — exécution", style="List Bullet")

    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for index, entete in enumerate(("Poste", "Quantité", "Montant")):
        cellule = table.cell(0, index)
        cellule.text = entete
        for run in cellule.paragraphs[0].runs:
            run.font.bold = True
            run.font.size = Pt(10)
    for index, ligne in enumerate(
        (("Formateur", "2", "12 000,00"), ("Salle", "1", "1 500,00")), start=1
    ):
        for colonne, valeur in enumerate(ligne):
            table.cell(index, colonne).text = valeur
    for ligne in table.rows:
        for cellule in ligne.cells:
            cellule.width = Mm(50)

    document.add_page_break()
    document.add_heading("Annexe", level=1)
    document.add_paragraph("Chiffres clés du lot.")
    document.add_picture(io.BytesIO(build_png(6, 4)), width=Mm(30))
    entete = document.sections[0].header.paragraphs[0]
    entete.text = "CARSO — Offre technique (proposition)"
    pied = document.sections[0].footer.paragraphs[0]
    pied.text = "Document proposé"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_reference_stylisee() -> bytes:
    """DOCX de référence portant une **identité visuelle** exploitable.

    Polices/tailles/couleurs de titres distinctes du style CARSO par défaut,
    marges spécifiques : c'est ce que ``analyze_document_reference`` doit savoir
    extraire sans copier un seul mot du contenu.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor

    document = docx.Document()
    section = document.sections[0]
    section.left_margin = Mm(30)
    section.right_margin = Mm(25)
    section.top_margin = Mm(28)
    section.bottom_margin = Mm(22)

    titre = document.add_heading("MODÈLE DE PROPOSITION", level=0)
    for run in titre.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0B, 0x2B, 0x4A)
        run.font.italic = False
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER

    section_titre = document.add_heading("1. Contexte", level=1)
    for run in section_titre.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x0B, 0x2B, 0x4A)

    corps = document.add_paragraph("Contenu de référence — ne doit jamais être recopié.")
    for run in corps.runs:
        run.font.name = "Georgia"
        run.font.size = Pt(10.5)

    table = document.add_table(rows=2, cols=2)
    table.style = "Light Grid Accent 1"
    table.cell(0, 0).text = "Rubrique"
    table.cell(0, 1).text = "Valeur"
    table.cell(1, 0).text = "Durée"
    table.cell(1, 1).text = "5 jours"

    for paragraphe in document.sections[0].header.paragraphs:
        paragraphe.text = "MODÈLE CLIENT — usage interne"
    for paragraphe in document.sections[0].footer.paragraphs:
        paragraphe.text = "Page 1"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf_stylise(pages: int = 2) -> bytes:
    """PDF multi-pages avec tailles de police distinctes (référence de style)."""
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(25, 25, 25)
    for index in range(pages):
        pdf.add_page()
        pdf.set_font("Times", "B", 18)
        pdf.cell(0, 10, f"Titre de page {index + 1}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Times", size=11)
        pdf.multi_cell(0, 6, "Corps de texte du document de reference.")
        pdf.set_font("Times", size=9)
        pdf.cell(0, 5, f"Pied de page {index + 1}", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def build_xlsx_riche() -> bytes:
    """XLSX réaliste : fusion, formules, styles, largeurs de colonnes."""
    classeur = Workbook()
    onglet = classeur.active
    onglet.title = "Bénéficiaires"
    onglet["A1"] = "Liste des bénéficiaires — lot 3"
    onglet["A1"].font = Font(bold=True, size=13)
    onglet.merge_cells("A1:D1")
    onglet.append(["Nom", "Prénom", "Présence", "Total"])
    for cellule in onglet["A2:D2"]:
        for c in cellule:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
    onglet.append(["Karimi", "Salma", 1, "=C3*10"])
    onglet.append(["Alaoui", "Youssef", 0, "=C4*10"])
    onglet["D5"] = "=SUM(D3:D4)"
    onglet.column_dimensions["A"].width = 22
    onglet.column_dimensions["B"].width = 18

    buffer = io.BytesIO()
    classeur.save(buffer)
    return buffer.getvalue()


def write(path: Path, contenu: bytes) -> Path:
    """Écrit des octets sur disque et renvoie le chemin."""
    path.write_bytes(contenu)
    return path
