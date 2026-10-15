"""Renderer XLSX — un ``DocumentSpec`` devient un classeur lisible (ADR 0007).

Les tests réinspectent le **fichier produit** avec openpyxl, la librairie de
production : ils vérifient le classeur réellement écrit (onglets, styles, gel,
filtre, fusions), pas l'appel de fonction.

Rappel du trou fermé par cet incrément : ``rendre(spec, format="xlsx")`` levait
« Format de rendu inconnu » (le format n'existait pas), et le seul chemin XLSX
disponible (``build_xlsx``) écrivait un classeur sans aucun style.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl
import pytest

from app.core.errors import GenerationError
from app.documents.generation import build_xlsx
from app.documents.renderers import FORMATS_RENDUS, extension_et_mime, rendre
from app.documents.renderers.xlsx import (
    MAX_CARACTERES_ONGLET,
    nom_onglet,
    rendre_tableur,
)
from app.documents.spec import (
    Citation,
    DocumentSpec,
    Image,
    Liste,
    Metadonnees,
    Paire,
    Paires,
    Paragraphe,
    Tableau,
    Titre,
)
from app.documents.styles import CARSO_DEFAUT, DocumentStyle, StyleTableau


def _classeur(octets: bytes) -> Any:
    return openpyxl.load_workbook(io.BytesIO(octets))


def _spec_document() -> DocumentSpec:
    """Document complet : métadonnées, texte, deux tableaux (dont un nommé)."""
    return DocumentSpec(
        metadata=Metadonnees(
            titre="Offre technique — Lot 3",
            sous_titre="Appel AAP-2026-07",
            organisation="CARSO",
            reference="OT-2026-014",
            auteur="AgentGenerateurOffre",
        ),
        blocs=[
            Titre(niveau=1, texte="Contexte"),
            Paragraphe(texte="Une phrase de contexte pour la lecture du classeur."),
            Liste(items=["Premier point", "Second point"], numerotee=True),
            Tableau(
                titre_tableau="Budget",
                entete=["Poste", "Montant"],
                lignes=[["Formation", "1200"], ["Logistique", "300"]],
                alignements=["gauche", "droite"],
            ),
            Tableau(
                entete=["Bénéficiaire", "Statut"],
                lignes=[["Awa", "Inscrit"]],
            ),
        ],
    )


# --- Format déclaré -----------------------------------------------------------


def test_format_xlsx_declare() -> None:
    """Le format est déclaré une seule fois (extension + type MIME)."""
    assert "xlsx" in FORMATS_RENDUS
    extension, mime = extension_et_mime("xlsx")
    assert extension == ".xlsx"
    assert mime == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_rendre_xlsx_ne_leve_plus_format_inconnu() -> None:
    """Régression : le rendu XLSX d'un document structuré échouait."""
    octets = rendre(_spec_document(), format="xlsx")
    assert octets[:2] == b"PK"  # un classeur est une archive ZIP


# --- Structure du classeur ----------------------------------------------------


def test_une_feuille_par_tableau() -> None:
    classeur = _classeur(rendre(_spec_document(), format="xlsx"))
    assert classeur.sheetnames[0] == "Offre technique — Lot 3"
    assert classeur.sheetnames[1:] == ["Budget", "Tableau 2"]


def test_feuille_document_porte_metadonnees_et_textes() -> None:
    classeur = _classeur(rendre(_spec_document(), format="xlsx"))
    feuille = classeur[classeur.sheetnames[0]]
    valeurs = [feuille.cell(row=index, column=1).value for index in range(1, 20)]
    assert "Offre technique — Lot 3" == valeurs[0]
    assert "Référence" in valeurs
    assert "OT-2026-014" in [feuille.cell(row=index, column=2).value for index in range(1, 20)]
    assert "Contexte" in valeurs
    assert "1. Premier point" in valeurs
    assert "2. Second point" in valeurs


def test_feuille_document_ignore_les_blocs_tableau() -> None:
    """Un tableau va sur sa feuille : la feuille document ne le duplique pas."""
    classeur = _classeur(rendre(_spec_document(), format="xlsx"))
    feuille = classeur[classeur.sheetnames[0]]
    textes = [
        cellule.value
        for rangee in feuille.iter_rows()
        for cellule in rangee
        if isinstance(cellule.value, str)
    ]
    assert "Poste" not in textes
    assert "Budget" not in textes  # le titre du tableau reste sur son onglet


# --- Style --------------------------------------------------------------------


def test_entete_stylise_et_bandes_alternees() -> None:
    """L'en-tête porte le style du design system, les lignes alternent.

    Layout d'un onglet tableau : ligne 1 = titre du tableau **s'il est déclaré**,
    puis la ligne d'en-tête, puis les données.
    """
    tokens = CARSO_DEFAUT.tableau
    classeur = _classeur(rendre(_spec_document(), format="xlsx"))
    feuille = classeur["Budget"]
    assert feuille["A1"].value == "Budget"  # titre du tableau, pas l'en-tête
    entete = feuille["A2"]
    assert entete.font.bold is tokens.entete_gras
    assert entete.font.color.rgb.endswith(tokens.entete_couleur.lstrip("#"))
    assert entete.fill.start_color.rgb.endswith(tokens.entete_fond.lstrip("#"))
    assert entete.border.left.style is not None  # bordures actives par défaut

    premiere_donnee, seconde_donnee = feuille["A3"], feuille["A4"]
    assert premiere_donnee.fill.start_color.rgb != seconde_donnee.fill.start_color.rgb
    assert seconde_donnee.fill.start_color.rgb.endswith(
        tokens.bande_fond.lstrip("#")
    )


def test_entete_en_premiere_ligne_sans_titre_de_tableau() -> None:
    """Sans titre de tableau déclaré, l'en-tête occupe bien la ligne 1."""
    feuille = _classeur(rendre(_spec_document(), format="xlsx"))["Tableau 2"]
    assert feuille["A1"].value == "Bénéficiaire"
    assert feuille["A1"].font.bold is True
    assert feuille.freeze_panes == "A2"


def test_style_du_document_est_respecte() -> None:
    """Un style explicite change le rendu : le renderer applique, il ne décide pas."""
    style = DocumentStyle(
        nom="sobre",
        tableau=StyleTableau(
            entete_fond="#000000", entete_couleur="#FFFFFF", bandes_alternees=False
        ),
    )
    spec = DocumentSpec(
        metadata=Metadonnees(titre="Sobre"),
        style=style,
        blocs=[
            Tableau(entete=["A"], lignes=[["1"], ["2"]]),
        ],
    )
    feuille = _classeur(rendre(spec, format="xlsx"))["Tableau 1"]
    assert feuille["A1"].fill.start_color.rgb == "FF000000"
    assert feuille["A2"].fill.start_color.rgb == "00000000"  # aucune bande


def test_alignements_de_colonne_appliques() -> None:
    feuille = _classeur(rendre(_spec_document(), format="xlsx"))["Budget"]
    assert feuille["A2"].alignment.horizontal == "left"
    assert feuille["B3"].alignment.horizontal == "right"


def test_fusions_reprises() -> None:
    """``fusion_colonnes: 'fin'`` devient une fusion réelle dans le classeur."""
    spec = DocumentSpec(
        metadata=Metadonnees(titre="Section pleine largeur"),
        blocs=[
            Tableau(
                entete=["A", "B"],
                lignes=[
                    [{"texte": "Section", "fusion_colonnes": "fin"}],
                    ["x", "y"],
                ],
            )
        ],
    )
    feuille = _classeur(rendre(spec, format="xlsx"))["Tableau 1"]
    # Ligne 1 : en-tête déclaré ; ligne 2 : la cellule de section fusionnée.
    assert "A2:B2" in {str(fusion) for fusion in feuille.merged_cells.ranges}


def test_cellule_qui_commence_par_egal_reste_du_texte() -> None:
    """Un contenu de document n'est jamais exécuté comme une formule."""
    spec = DocumentSpec(
        metadata=Metadonnees(titre="Formules"),
        blocs=[Tableau(entete=["Expression"], lignes=[["=1+1"]])],
    )
    feuille = _classeur(rendre(spec, format="xlsx"))["Tableau 1"]
    cellule = feuille["A2"]
    assert cellule.value == "=1+1"
    assert cellule.data_type == "s"


# --- Lisibilité ---------------------------------------------------------------


def test_gele_filtre_et_titre_impression() -> None:
    """En-tête gelé sous la ligne d'en-tête, filtre sur le tableau, répétition."""
    feuille = _classeur(rendre(_spec_document(), format="xlsx"))["Budget"]
    assert feuille.freeze_panes == "A3"
    assert feuille.auto_filter.ref == "A2:B4"
    assert feuille.print_title_rows == "$2:$2"  # normalisé par openpyxl
    assert feuille.page_setup.orientation == "landscape"


def test_largeurs_declarees_converties_en_caracteres() -> None:
    """``largeurs_mm`` (unité du document) devient une largeur de colonne."""
    spec = DocumentSpec(
        metadata=Metadonnees(titre="Largeurs"),
        blocs=[
            Tableau(
                entete=["A", "B"],
                lignes=[["x", "y"]],
                largeurs_mm=[37.0, 18.5],
            )
        ],
    )
    feuille = _classeur(rendre(spec, format="xlsx"))["Tableau 1"]
    assert feuille.column_dimensions["A"].width == pytest.approx(22.0, abs=0.01)
    assert feuille.column_dimensions["B"].width == pytest.approx(12.0, abs=0.01)


def test_image_signalee_sans_etre_perdue() -> None:
    """Un classeur n'intègre pas les images : leur place est écrite, pas effacée."""
    spec = DocumentSpec(
        metadata=Metadonnees(titre="Avec image"),
        blocs=[Image(chemin="missions/1/plan.png", legende="Plan de salle")],
    )
    feuille = _classeur(rendre(spec, format="xlsx")).worksheets[0]
    textes = [
        cellule.value
        for rangee in feuille.iter_rows()
        for cellule in rangee
        if isinstance(cellule.value, str)
    ]
    assert any(
        "image non reportée" in texte and "Plan de salle" in texte for texte in textes
    )


def test_citation_et_paires_ecrites() -> None:
    spec = DocumentSpec(
        metadata=Metadonnees(titre="Textes"),
        blocs=[
            Citation(texte="Décider, c'est choisir.", attribution="A. Carso"),
            Paires(paires=[Paire(libelle="Durée", valeur="3 jours")]),
        ],
    )
    feuille = _classeur(rendre(spec, format="xlsx")).worksheets[0]
    textes = [
        cellule.value
        for rangee in feuille.iter_rows()
        for cellule in rangee
        if isinstance(cellule.value, str)
    ]
    assert "« Décider, c'est choisir. » — A. Carso" in textes
    assert "Durée" in textes and "3 jours" in textes


# --- Noms d'onglets -----------------------------------------------------------


def test_nom_onglet_limite_et_unicite() -> None:
    pris: set[str] = set()
    long = nom_onglet("B" * 60, defaut="Tableau", pris=pris)
    assert len(long) == MAX_CARACTERES_ONGLET
    meme = nom_onglet("B" * 60, defaut="Tableau", pris=pris)
    assert meme != long
    assert len(meme) <= MAX_CARACTERES_ONGLET
    # Les caractères interdits par Excel sont neutralisés.
    assert "[" not in nom_onglet("Budget [T1]/x", defaut="T", pris=set())
    assert nom_onglet(None, defaut="Tableau 3", pris=set()) == "Tableau 3"


def test_deux_tableaux_homonymes_ont_deux_onglets_distincts() -> None:
    spec = DocumentSpec(
        metadata=Metadonnees(titre="Homonymes"),
        blocs=[
            Tableau(titre_tableau="Budget", entete=["A"], lignes=[["1"]]),
            Tableau(titre_tableau="Budget", entete=["A"], lignes=[["2"]]),
        ],
    )
    classeur = _classeur(rendre(spec, format="xlsx"))
    assert classeur.sheetnames[1] == "Budget"
    assert classeur.sheetnames[2].startswith("Budget")
    assert len(set(classeur.sheetnames)) == 3


# --- Chemin historique : build_xlsx -------------------------------------------


def test_build_xlsx_est_style() -> None:
    """La fiche de présence passe par le même style que tout autre classeur."""
    tokens = CARSO_DEFAUT.tableau
    octets = build_xlsx(
        titre_feuille="Presence",
        lignes=[["Nom", "Présence"], ["Awa", "Présent"], ["Bilal", "Absent"]],
    )
    classeur = _classeur(octets)
    assert classeur.sheetnames == ["Presence"]
    feuille = classeur["Presence"]
    entete = feuille["A1"]
    assert entete.font.bold is tokens.entete_gras
    assert entete.fill.start_color.rgb.endswith(tokens.entete_fond.lstrip("#"))
    assert feuille.freeze_panes == "A2"
    assert feuille.auto_filter.ref == "A1:B3"
    assert feuille["A3"].fill.start_color.rgb.endswith(tokens.bande_fond.lstrip("#"))


def test_build_xlsx_sans_ligne_reste_un_classeur_valide() -> None:
    classeur = _classeur(build_xlsx(titre_feuille="Vide", lignes=[]))
    assert classeur.sheetnames == ["Vide"]


def test_rendre_tableur_une_seule_feuille() -> None:
    octets = rendre_tableur(
        titre_feuille="Liste des bénéficiaires",
        lignes=[["Nom", "Statut"], ["Awa", "Inscrit"]],
        style=CARSO_DEFAUT,
    )
    classeur = _classeur(octets)
    assert classeur.sheetnames == ["Liste des bénéficiaires"]


def test_rendre_tableur_librairie_absente(monkeypatch: pytest.MonkeyPatch) -> None:
    """Sans openpyxl : erreur explicite, jamais un fichier vide."""
    from app.core import errors

    def _leve(nom: str) -> object:
        raise errors.GenerationError(f"Librairie absente : {nom}")

    monkeypatch.setattr(
        "app.documents.extraction.base.require_library", _leve, raising=True
    )
    with pytest.raises(GenerationError):
        rendre_tableur(
            titre_feuille="X", lignes=[["A"]], style=CARSO_DEFAUT
        )
