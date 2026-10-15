"""Tests — tableaux du ``DocumentSpec`` (en-tête, fusions, trames) et import XLSX.

Ce fichier verrouille deux défauts constatés sur un document réel :

1. un tableau **sans en-tête déclaré** recevait une ligne d'en-tête tramée
   **vide**, et les vrais titres de colonnes tombaient dans la ligne suivante ;
2. une structure complexe (en-tête à étages, ligne de section pleine largeur,
   cellules tramées) était impossible à exprimer.

Le rendu est vérifié dans les **octets réels** produits (document.xml du DOCX,
HTML), pas sur une intention.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError as PydanticValidationError

from app.core.errors import ValidationError
from app.documents.renderers import rendre
from app.documents.renderers.html import rendre_html
from app.documents.spec import Cellule, DocumentSpec, Metadonnees, Tableau
from app.documents.tableau_xlsx import import_tableau_xlsx, tableau_depuis_feuille


def _docx_xml(blocs: list[Tableau]) -> str:
    """XML du corps DOCX pour une liste de tableaux (octets réels)."""
    octets = rendre(
        DocumentSpec(metadata=Metadonnees(titre="Tableaux"), blocs=blocs), format="docx"
    )
    return zipfile.ZipFile(io.BytesIO(octets)).read("word/document.xml").decode("utf-8")


def _lignes(xml: str) -> list[str]:
    return re.findall(r"<w:tr[ >].*?</w:tr>", xml, re.S)


def _textes(ligne: str) -> list[str]:
    return re.findall(r"<w:t[^>]*>([^<]*)</w:t>", ligne)


class TestEnteteDeTableau:
    """Un tableau sans en-tête déclaré n'a pas de bandeau vide."""

    def test_sans_entete_aucune_ligne_tramee(self) -> None:
        xml = _docx_xml([Tableau(lignes=[["Période", "Activités"], ["Mois 1", "Préparation"]])])
        lignes = _lignes(xml)

        assert len(lignes) == 2
        assert _textes(lignes[0]) == ["Période", "Activités"]
        assert "1F4E79" not in lignes[0]

    def test_entete_declare_trame_et_texte(self) -> None:
        xml = _docx_xml(
            [Tableau(entete=["Période", "Activités"], lignes=[["Mois 1", "Préparation"]])]
        )
        lignes = _lignes(xml)

        assert _textes(lignes[0]) == ["Période", "Activités"]
        assert "1F4E79" in lignes[0]
        assert "1F4E79" not in lignes[1]

    def test_entete_et_entetes_exclusifs(self) -> None:
        """Deux façon d'écrire l'en-tête : en donner deux est refusé à la construction."""
        with pytest.raises(PydanticValidationError, match="pas les deux"):
            Tableau(entete=["a"], entetes=[["b"]], lignes=[["x"]])


class TestStructuresComplexes:
    """En-tête à étages, fusions verticales, ligne de section, trames."""

    def _tableau(self) -> Tableau:
        return Tableau(
            entetes=[
                [
                    Cellule(texte="Produits attendus", fusion_lignes=2),
                    Cellule(texte="Activités", fusion_lignes=2),
                    Cellule(texte="2026", fusion_colonnes=2),
                ],
                [Cellule(texte="1"), Cellule(texte="2")],
            ],
            lignes=[
                ["ACT4", "Suivi", "3", ""],
                [Cellule(texte="Produit 3 : incubés", fusion_colonnes="fin", fond="#E2EFDA")],
                [
                    Cellule(texte="TOTAL", gras=True),
                    Cellule(texte="66 jours", fusion_colonnes="fin", fond="#FCE4D6"),
                ],
            ],
        )

    def test_grille_place_les_fusions(self) -> None:
        tableau = self._tableau()

        assert tableau.largeur_grille() == 4
        assert tableau.nb_lignes_entete() == 2
        premiere = tableau.grille()[0]
        assert [(p.cellule.texte, p.colonne, p.colonnes, p.lignes) for p in premiere] == [
            ("Produits attendus", 0, 1, 2),
            ("Activités", 1, 1, 2),
            ("2026", 2, 2, 1),
        ]
        # La ligne de section couvre toute la largeur.
        section = tableau.grille()[3]
        assert [(p.colonne, p.colonnes) for p in section] == [(0, 4)]

    def test_grille_refuse_une_fusion_hors_grille(self) -> None:
        """Une cellule qui dépasse la largeur du tableau est refusée, pas tronquée."""
        bancal = Tableau(
            lignes=[
                [Cellule(texte="A", fusion_lignes=2), "B"],
                ["C", "D"],
            ]
        )
        with pytest.raises(ValidationError, match="au-delà des 2 colonnes"):
            bancal.grille()

    def test_grille_refuse_une_ligne_incomplete(self) -> None:
        bancal = Tableau(entete=["a", "b"], lignes=[["x"]])
        with pytest.raises(ValidationError, match="non couverte"):
            bancal.grille()

    def test_largeur_indeterminable_si_tout_en_fin(self) -> None:
        with pytest.raises(ValidationError, match="impossible de déterminer"):
            Tableau(lignes=[[Cellule(texte="a", fusion_colonnes="fin")]]).largeur_grille()

    def test_docx_fusions_et_trames(self) -> None:
        xml = _docx_xml([self._tableau()])
        lignes = _lignes(xml)

        assert len(lignes) == 5
        # En-tête à étages : fusion verticale sur la première ligne, horizontale sur 2026.
        assert len(re.findall(r"<w:vMerge", lignes[0])) == 2
        assert 'w:gridSpan w:val="2"' in lignes[0]
        # Ligne de section : une seule cellule, sur toute la largeur, tramée.
        assert _textes(lignes[3]) == ["Produit 3 : incubés"]
        assert "E2EFDA" in lignes[3]
        assert 'w:gridSpan w:val="4"' in lignes[3]

    def test_html_colspan_rowspan_et_trames(self) -> None:
        html = rendre_html(
            DocumentSpec(metadata=Metadonnees(titre="T"), blocs=[self._tableau()])
        ).decode("utf-8")

        assert '<th rowspan="2">' in html
        assert '<th colspan="2">' in html
        assert 'colspan="4" style="background:#E2EFDA"' in html
        assert 'style="font-weight:bold"' in html

    def test_pdf_rend_toutes_les_cellules(self) -> None:
        from pypdf import PdfReader

        from app.documents.spec import Couverture, MiseEnPage

        spec = DocumentSpec(
            metadata=Metadonnees(titre="T"),
            mise_en_page=MiseEnPage(couverture=Couverture(active=False)),
            blocs=[self._tableau()],
        )
        texte = PdfReader(io.BytesIO(rendre(spec, format="pdf"))).pages[0].extract_text() or ""

        for attendu in ("Produits attendus", "2026", "ACT4", "Produit 3", "TOTAL", "66 jours"):
            assert attendu in texte


class TestImportXlsx:
    """Un tableau déjà saisi dans un classeur se reprend tel quel."""

    def _classeur(self, chemin: Path) -> Path:
        classeur = Workbook()
        feuille = classeur.active
        feuille.title = "Annexe C"
        feuille["A1"] = "Produits attendus"
        feuille["B1"] = "Activités"
        feuille["C1"] = "2026"
        feuille["C1"].font = Font(bold=True)
        feuille["C1"].fill = PatternFill("solid", start_color="FF1F4E79")
        feuille["A2"] = "ACT4"
        feuille["B2"] = "Réaliser le suivi"
        feuille["C2"] = 3
        feuille["A3"] = "TOTAL"
        feuille["B3"] = "66 jours"
        feuille.merge_cells("B3:C3")
        feuille["B3"].alignment = Alignment(horizontal="center")
        feuille.column_dimensions["A"].width = 24
        classeur.save(chemin)
        return chemin

    def test_import_conserve_fusions_trames_et_largeurs(self, tmp_path: Path) -> None:
        tableau = import_tableau_xlsx(
            self._classeur(tmp_path / "annexe.xlsx"),
            feuille="Annexe C",
            ligne_entete=1,
        )

        assert tableau.nb_lignes_entete() == 1
        assert tableau.largeur_grille() == 3
        assert [cellule.texte for cellule in tableau.lignes_entete()[0]] == [
            "Produits attendus",
            "Activités",
            "2026",
        ]
        entete_2026 = tableau.lignes_entete()[0][2]
        assert entete_2026.fond == "#1F4E79"
        assert entete_2026.gras is True
        # La fusion Excel devient une fusion du spec, et le style est conservé.
        total = tableau.lignes[1][1]
        assert total.texte == "66 jours"
        assert total.fusion_colonnes == 2
        assert total.alignement == "centre"
        assert tableau.largeurs_mm and tableau.largeurs_mm[0] > tableau.largeurs_mm[1]

    def test_import_avec_titre_fusionne_en_premiere_ligne(self, tmp_path: Path) -> None:
        """Une première ligne fusionnée ne fait plus échouer l'import complet.

        Les classeurs réels (budget TEFIALA) fusionnent leur titre sur toute la
        largeur : la lettre de colonne lue sur la cellule rendait un
        ``MergedCell``, et ``import_tableau_xlsx`` levait sur un assistant
        ``column_letter`` — aucune ligne n'était importée (constat 04/10).
        """
        chemin = tmp_path / "budget.xlsx"
        classeur = Workbook()
        feuille = classeur.active
        feuille["A1"] = "Budget détaillé"
        feuille.merge_cells("A1:C1")
        feuille["A2"], feuille["B2"], feuille["C2"] = "Poste", "Unité", "Montant"
        feuille["A3"], feuille["B3"], feuille["C3"] = "Formation", "jour", 12
        feuille.column_dimensions["A"].width = 30
        classeur.save(chemin)

        tableau = import_tableau_xlsx(chemin, ligne_entete=0)

        assert tableau.largeur_grille() == 3
        assert tableau.largeurs_mm[0] > 0
        assert "Budget détaillé" in tableau.lignes[0][0].texte

    def test_import_sans_entete(self, tmp_path: Path) -> None:
        tableau = import_tableau_xlsx(
            self._classeur(tmp_path / "annexe.xlsx"), ligne_entete=0
        )
        assert tableau.nb_lignes_entete() == 0
        assert tableau.lignes[0][0].texte == "Produits attendus"

    def test_import_feuille_inconnue(self, tmp_path: Path) -> None:
        with pytest.raises(ValidationError, match="Feuille inconnue"):
            import_tableau_xlsx(self._classeur(tmp_path / "annexe.xlsx"), feuille="Absente")

    def test_import_du_tableau_se_rend(self, tmp_path: Path) -> None:
        tableau = import_tableau_xlsx(self._classeur(tmp_path / "annexe.xlsx"))
        xml = _docx_xml([tableau])

        assert "ACT4" in xml
        assert "1F4E79" in xml
        assert 'w:gridSpan w:val="2"' in xml

    def test_cellule_vide_reste_vide(self) -> None:
        classeur = Workbook()
        feuille = classeur.active
        feuille["A1"] = "Zone"
        feuille["B1"] = "Responsable"
        feuille["A2"] = "Rabat"
        tableau = tableau_depuis_feuille(feuille, ligne_entete=1)

        assert tableau.lignes[0][1].texte == ""
        assert tableau.largeur_grille() == 2

    def test_plage_sans_donnee_refusee(self) -> None:
        classeur = Workbook()
        feuille = classeur.active
        feuille["A1"] = "Zone"
        with pytest.raises(ValidationError, match="Aucune ligne de données"):
            tableau_depuis_feuille(feuille, ligne_entete=1)
