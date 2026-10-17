"""Classeur Excel professionnel (``format='xlsx'``) + visibilité « genere_pour ».

Trois points de contrôle :

1. **moteur** (``xlsx_renderer``) : chaque ``table.data`` devient une feuille
   stylée nommée d'après le titre qui la précède, les ``table.kpi`` une feuille
   « Chiffres clés », en-tête figé et filtre automatique ; sans tableau, refus
   explicite (jamais un .xlsx vide livré) ;
2. **rendu** (``rendre_professionnel(format='xlsx')``) : octets OOXML valides
   (+ PDF jumeau quand un navigateur est présent) et avertissement explicite si
   ``graphiques`` est passé (sans effet sur un classeur) ;
3. **visibilité** (incrément 74) : un document créé sous périmètre restreint
   (généré par l'agent, ou déposé) reste visible de **son** compte, et de lui
   seul — même sans ancre métier.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterator
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.application.dto import DocumentUploadInput
from app.application.perimetre import Perimetre
from app.application.services import DocumentService
from app.core.errors import NotFoundError, ValidationError
from app.documents.professionnel import rendre_professionnel
from app.documents.storage import LocalDocumentStorage
from app.domain.base import Base

HTML_TABLEAUX = """
<h1>Budget prévisionnel</h1>
<table class="data">
  <thead><tr><th>Poste</th><th class="num">Montant</th></tr></thead>
  <tbody>
    <tr><td>Formation</td><td class="num">4 800</td></tr>
    <tr><td>Logistique</td><td class="num">1 200</td></tr>
  </tbody>
</table>
<h2>Effectifs</h2>
<table class="data">
  <thead><tr><th>Session</th><th class="num">Inscrits</th></tr></thead>
  <tbody><tr><td>Lot 1</td><td class="num">12</td></tr></tbody>
</table>
<table class="kpi"><tr>
  <td><span class="kpi-value">+24 %</span><span class="kpi-label">Croissance</span></td>
</tr></table>
"""


# --- Moteur --------------------------------------------------------------------


def test_xlsx_une_feuille_par_tableau(tmp_path: Path) -> None:
    chemin = rendre_professionnel(HTML_TABLEAUX, "Budget", format="xlsx").octets
    classeur = load_workbook(io.BytesIO(chemin))

    assert classeur.sheetnames[:2] == ["Budget prévisionnel", "Effectifs"]
    assert classeur.sheetnames[-1] == "Chiffres clés"
    feuille = classeur["Budget prévisionnel"]
    assert feuille["A2"].value == "Poste"
    assert feuille["B3"].value == "4 800"
    assert feuille.freeze_panes == "A3"
    assert feuille.auto_filter.ref is not None
    # En-tête stylé : fond primaire, blanc gras (thème corporate).
    assert feuille["A2"].fill.start_color.rgb == "001E3A5F"
    assert feuille["A2"].font.bold is True


def test_xlsx_refuse_sans_tableau(tmp_path: Path) -> None:
    with pytest.raises(ValidationError) as excinfo:
        rendre_professionnel("<h1>Rapport</h1><p>Du texte, aucun tableau.</p>",
                             "Rapport", format="xlsx")
    assert "tableau" in str(excinfo.value.message).lower()


def test_xlsx_annonce_que_les_graphiques_sont_sans_effet(tmp_path: Path) -> None:
    rendu = rendre_professionnel(
        HTML_TABLEAUX,
        "Budget",
        format="xlsx",
        graphiques={"repartition": {"type": "bar", "labels": ["T1"], "series": {"A": [1]}}},
    )
    assert any("sans effet sur un classeur" in a for a in rendu.avertissements)


def test_xlsx_ne_fabrique_pas_un_docx(tmp_path: Path) -> None:
    rendu = rendre_professionnel(HTML_TABLEAUX, "Budget", format="xlsx")
    assert rendu.format == "xlsx"
    assert rendu.octets[:2] == b"PK"

    with zipfile.ZipFile(io.BytesIO(rendu.octets)) as archive:
        assert "xl/workbook.xml" in archive.namelist()
        assert "word/document.xml" not in archive.namelist()


# --- Visibilité « genere_pour » (incrément 74) ---------------------------------


@pytest.fixture
def session() -> Iterator[Session]:
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as active:
        yield active
    engine.dispose()


@pytest.fixture
def storage(tmp_path: Path) -> LocalDocumentStorage:
    return LocalDocumentStorage(tmp_path / "carso", max_bytes=5 * 1024 * 1024)


def test_document_genere_visible_de_son_compte_seulement(
    session: Session, storage: LocalDocumentStorage
) -> None:
    """Le formateur retrouve ce qu'il fait générer par l'agent, sans ancre."""
    utilisateur = uuid4()
    restreint = DocumentService(
        session,
        storage=storage,
        actor_id="agent_assistant_formateur",
        perimetre=Perimetre(role="formateur", utilisateur_id=utilisateur, est_complet=False),
    )
    document = restreint.enregistrer_document(
        DocumentUploadInput(
            type_document="rapport",
            nom="note_agent.txt",
            stream=io.BytesIO("Note générée par l'agent.".encode("utf-8")),
            created_by="agent_assistant_formateur",
        )
    )

    # La trace du demandeur est écrite, et lui seul y a accès.
    assert document.doc_metadata is not None
    assert document.doc_metadata["genere_pour"] == str(utilisateur)
    assert restreint.obtenir(document.id).id == document.id
    assert [d.id for d in restreint.lister_tous()] == [document.id]

    autre = DocumentService(
        session,
        storage=storage,
        actor_id="observateur",
        perimetre=Perimetre(role="formateur", utilisateur_id=uuid4(), est_complet=False),
    )
    with pytest.raises(NotFoundError):
        autre.obtenir(document.id)
    assert autre.lister_tous() == []
